#!/usr/bin/env python3
# GrepAll - Offline OCR for PDFs and images, with Arabic support
# Copyright (C) 2026 A.T.Grep
# Licensed under the GNU Affero General Public License v3.0. See LICENSE.
"""GrepAll - single-file desktop app for Windows and Linux.

Creates searchable PDFs (original pages kept unchanged, invisible OCR text layer added)
and optionally converts them to Markdown.

Dependencies: python -m pip install pymupdf pymupdf4llm
Tesseract language data (eng.traineddata, ara.traineddata) is also required, either
installed with Tesseract or placed in a "tessdata" folder next to this script or the .exe.
"""
from __future__ import annotations

import json
import math
import os
import platform
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.request
from collections import Counter
import threading
import time
import tkinter as tk
import tkinter.font as tkfont
import webbrowser
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from urllib.parse import quote

try:
    import pymupdf
except ImportError:
    try:
        import fitz as pymupdf
    except ImportError as exc:
        raise SystemExit("PyMuPDF is required: python -m pip install pymupdf") from exc

APP_NAME = "GrepAll"
APP_TAGLINE = "Offline OCR for PDFs and images, with Arabic support"
APP_VERSION = "1.6.0"
APP_CREDIT = "A.T.Grep"
FEEDBACK_EMAIL = "dev.oasis006@passmail.net"
REPOSITORY = "swaidi/grepall"
RELEASES_URL = f"https://github.com/{REPOSITORY}/releases/latest"
UPDATE_CHECK_INTERVAL = 24 * 60 * 60  # seconds
ICON_FILE = "app_icon.png"


def resource_path(name: str) -> Path:
    """Locate a bundled file, both when running as a script and as a PyInstaller .exe."""
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / name

MIN_TEXT_CHARS = 25
MAX_RENDER_PIXELS = 60_000_000  # caps memory use on large drawings (A0/A1)
MAILTO_SAFE_LENGTH = 1800       # longer mailto links are truncated by some mail apps
LANGUAGES = {"Arabic + English": "ara+eng", "English": "eng", "Arabic": "ara"}
RESOLUTIONS = {"Compact (200 dpi)": 200, "Standard (300 dpi)": 300, "Small print (400 dpi)": 400}
ARABIC = re.compile(r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]")
LTR_CHARS = re.compile(r"[A-Za-z0-9\u0660-\u0669]")
MD_PREFIX = re.compile(r"^(\s*(?:#{1,6}\s+|[-*+]\s+|>\s+)?)(.*)$")
IMAGE_SUFFIXES = (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp")
OFFICE_SUFFIXES = (".docx", ".xlsx", ".xlsm", ".pptx")  # converted to Markdown only
INPUT_SUFFIXES = (".pdf",) + IMAGE_SUFFIXES + OFFICE_SUFFIXES
OUTPUT_STEMS = ("_ocr", "_small", "_redacted", "_tables", ".partial")  # GrepAll's own results
IMAGE_PAGE_LONG_SIDE = 842  # images are placed on A4-sized pages (842 pt = 297 mm)
ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫٬٪", "01234567890123456789.,%")
ARABIC_DIACRITICS = re.compile(r"[\u064B-\u065F\u0670]")
ALEF_FORMS = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا"})
TATWEEL = "\u0640"

DEFAULT_SETTINGS = {
    "language": "Arabic + English",
    "quality": "Standard (300 dpi)",
    "output_dir": "",
    "combine_images": False,
    "markers": True,
    "strip_headers": False,
    "arabic_digits": True,
    "arabic_tatweel": True,
    "arabic_diacritics": True,
    "arabic_alef": True,
    "redact": False,
    "redact_pdf": True,
    "redact_terms": "",
    "redact_labels": "",
    "check_updates": True,
    "last_update_check": 0.0,
    "theme": "dark",
    "fix_rotation": True,
    "save_searchable": False,
    "reduce_level": "Balanced",
    "reduce_gray": False,
    "library_auto": True,
    "last_page": "ocr",
    "log_open": False,
    "window_geometry": "1180x800",
}


def settings_path() -> Path:
    """Per-user settings file: %APPDATA%\\GrepAll on Windows, ~/.config/grepall on Linux."""
    return _settings_base() / ("GrepAll" if sys.platform.startswith("win") else "grepall") / "settings.json"


def _settings_base() -> Path:
    if sys.platform.startswith("win"):
        return Path(os.environ.get("APPDATA") or Path.home() / "AppData" / "Roaming")
    return Path(os.environ.get("XDG_CONFIG_HOME") or Path.home() / ".config")


def _migrate_old_settings() -> None:
    """Copy settings saved by versions released as "PDF OCR to Markdown", once."""
    new = settings_path()
    old = _settings_base() / ("PDF-OCR" if sys.platform.startswith("win") else "pdf-ocr") / "settings.json"
    if new.exists() or not old.exists():
        return
    try:
        new.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(old, new)
    except OSError:
        pass


def load_settings() -> dict:
    _migrate_old_settings()
    settings = dict(DEFAULT_SETTINGS)
    try:
        stored = json.loads(settings_path().read_text(encoding="utf-8"))
    except Exception:
        return settings
    for key, default in DEFAULT_SETTINGS.items():
        value = stored.get(key)
        if isinstance(value, type(default)):
            settings[key] = value
    if settings["language"] not in LANGUAGES:
        settings["language"] = DEFAULT_SETTINGS["language"]
    if settings["quality"] not in RESOLUTIONS:
        settings["quality"] = DEFAULT_SETTINGS["quality"]
    if not re.fullmatch(r"\d{3,4}x\d{3,4}", settings["window_geometry"]):
        settings["window_geometry"] = DEFAULT_SETTINGS["window_geometry"]
    if settings["theme"] not in ("dark", "light"):
        settings["theme"] = "dark"
    if settings["reduce_level"] not in REDUCE_LEVELS:
        settings["reduce_level"] = "Balanced"
    return settings


def save_settings(settings: dict) -> None:
    path = settings_path()
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        temp = path.with_name(path.name + ".partial")
        temp.write_text(json.dumps(settings, indent=2, ensure_ascii=False), encoding="utf-8")
        os.replace(temp, path)
    except OSError:
        pass  # settings are a convenience; never block the app


def is_image(path: Path) -> bool:
    return path.suffix.lower() in IMAGE_SUFFIXES


def images_to_pdf(images: list[Path], dst: Path) -> int:
    """Combine image files into one PDF with A4-sized pages. Multi-page TIFFs add one page per frame.
    JPEG data is kept as is, and phone-photo rotation (EXIF) is respected. Returns the page count."""
    output = pymupdf.open()
    try:
        for image in images:
            with pymupdf.open(image) as image_doc:
                pdf_bytes = image_doc.convert_to_pdf()
            with pymupdf.open("pdf", pdf_bytes) as image_pdf:
                for index, page in enumerate(image_pdf):
                    scale = IMAGE_PAGE_LONG_SIDE / max(page.rect.width, page.rect.height)
                    new_page = output.new_page(width=page.rect.width * scale, height=page.rect.height * scale)
                    new_page.show_pdf_page(new_page.rect, image_pdf, index)
        if output.page_count == 0:
            raise RuntimeError("The image contains no pages.")
        output.save(dst, garbage=4, deflate=True)
        return output.page_count
    finally:
        output.close()


def clean_arabic(text: str, digits: bool, tatweel: bool, diacritics: bool, alef: bool) -> str:
    """Optional Arabic normalization for the Markdown output."""
    if digits:
        text = text.translate(ARABIC_DIGITS)
    if tatweel:
        text = text.replace(TATWEEL, "")
    if diacritics:
        text = ARABIC_DIACRITICS.sub("", text)
    if alef:
        text = text.translate(ALEF_FORMS)
    return text


# =============================================================== engine
def find_tessdata(explicit: str | None = None) -> str:
    candidates: list[str] = []
    if explicit:
        candidates.append(explicit)
    if getattr(sys, "frozen", False):
        candidates.extend((str(Path(getattr(sys, "_MEIPASS", "")) / "tessdata"),
                           str(Path(sys.executable).parent / "tessdata")))
    candidates.append(str(Path(__file__).resolve().parent / "tessdata"))
    if os.environ.get("TESSDATA_PREFIX"):
        candidates.append(os.environ["TESSDATA_PREFIX"])
    try:
        found = pymupdf.get_tessdata()
        if found:
            candidates.append(found)
    except Exception:
        pass
    candidates.extend((
        "/usr/share/tesseract-ocr/5/tessdata",
        "/usr/share/tesseract-ocr/4.00/tessdata",
        "/usr/share/tessdata",
        "/usr/local/share/tessdata",
        "/opt/homebrew/share/tessdata",
        r"C:\Program Files\Tesseract-OCR\tessdata",
        str(Path.home() / "AppData/Local/Programs/Tesseract-OCR/tessdata"),
    ))
    for folder in candidates:
        path = Path(folder) if folder else None
        if path and path.is_dir() and any(path.glob("*.traineddata")):
            return str(path)
    raise FileNotFoundError("Tesseract language data was not found. "
                            "Install Tesseract, or put a tessdata folder next to this app.")


def check_languages(tessdata: str, lang: str) -> None:
    missing = [item for item in lang.split("+") if not (Path(tessdata) / f"{item}.traineddata").is_file()]
    if missing:
        raise FileNotFoundError(f"Missing language data: {', '.join(missing)}\nExpected in: {tessdata}")


def parse_pages(spec: str | None, count: int) -> list[int]:
    """Convert a 1-based page selection such as '1-5, 9' to zero-based indexes."""
    if not spec:
        return list(range(count))
    result: set[int] = set()
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            first, last = part.split("-", 1)
            start, end = int(first), int(last) if last else count
        else:
            start = end = int(part)
        if start > end:
            start, end = end, start
        result.update(page - 1 for page in range(start, end + 1) if 1 <= page <= count)
    if not result:
        raise ValueError(f"No valid pages in '{spec}'. This PDF has {count} pages.")
    return sorted(result)


def strip_images(page, tiny_only: bool = False) -> int:
    """Remove image drawing commands from a page and the forms it uses.
    tiny_only removes only placeholder images of a few pixels. Returns the count removed."""
    doc = page.parent
    page.clean_contents()
    removed = 0
    for image in page.get_images(full=True):
        width, height, name, referencer = image[2], image[3], image[7], image[9]
        if tiny_only and width * height > 4:
            continue
        stream_xref = referencer or page.get_contents()[0]
        stream = doc.xref_stream(stream_xref)
        pattern = re.compile(rb"/" + re.escape(name.encode()) + rb"\s+Do\b")
        updated, count = pattern.subn(b"", stream)
        if count:
            doc.update_stream(stream_xref, updated)
            removed += count
        # Drop the resource entry too, so the image is not copied into the output file.
        doc.xref_set_key(referencer or page.xref, f"Resources/XObject/{name}", "null")
    return removed


class Cancelled(Exception):
    pass


def safe_dpi(page, dpi: int) -> int:
    """Lower the render resolution for very large pages so memory use stays bounded."""
    pixels = page.rect.width * page.rect.height * (dpi / 72) ** 2
    if pixels <= MAX_RENDER_PIXELS:
        return dpi
    return max(72, int(72 * math.sqrt(MAX_RENDER_PIXELS / (page.rect.width * page.rect.height))))


# Common words used to tell which way a page reads (sideways or upside-down scans read as gibberish).
_COMMON_WORDS = frozenset("""
the and of to in for with on at by from is are be this that as or an no not it its we our your you all any
date ref page total contract contractor works project shall will per item qty quantity rate amount price sar
subject dear sir regards letter notice delay supply materials site order purchase details number description
unit value tax vat invoice payment period days within accordance agreement scope work approved approval
""".split()) | frozenset("""
في من على إلى الى أن ان عن مع هذا هذه التي الذي رقم تاريخ بتاريخ المشروع العقد المقاول السعر الكمية ريال
الموقع خلال يوم أمر امر شركة المواد جميع بتوريد يلتزم المورد وذلك او أو قبل بعد كما بين حسب لا ما تم وفق
المبلغ القيمة الضريبة الدفع الأعمال الاعمال البند الوحدة الإجمالي الاجمالي المملكة العربية السعودية
""".split())


def reading_score(text: str) -> int:
    """How many common words a text contains; high for correctly oriented pages."""
    return sum(1 for word in re.findall(r"[A-Za-z\u0600-\u06FF]+", text.lower()) if word in _COMMON_WORDS)


def _ocr_bytes(page, angle: int, dpi: int, lang: str, tessdata: str) -> bytes:
    """OCR the page rendered turned by `angle` degrees; returns a one-page PDF with the text layer."""
    matrix = pymupdf.Matrix(dpi / 72, dpi / 72).prerotate(angle)
    pixmap = page.get_pixmap(matrix=matrix)
    pixmap.set_dpi(dpi, dpi)
    return pixmap.pdfocr_tobytes(compress=True, language=lang, tessdata=tessdata)


def detect_reading_angle(page, first_text: str, lang: str, tessdata: str) -> int | None:
    """Return the turn (0, 90, 180 or 270) that makes the page readable, or None when the page has
    too little text to tell (e.g. drawings). Pages that already read well cost nothing extra;
    others are tried quickly at low resolution in each direction."""
    base = reading_score(first_text)
    if base >= 3:
        return 0
    best_angle, best_score = 0, base
    for angle in (90, 270, 180):
        with pymupdf.open("pdf", _ocr_bytes(page, angle, 110, lang, tessdata)) as trial:
            score = reading_score(trial[0].get_text())
        if score > best_score:
            best_angle, best_score = angle, score
    return best_angle if best_score >= max(3, 2 * base + 1) else None


def ocr_pdf(src: Path, dst: Path, pages: list[int], lang: str, dpi: int,
            tessdata: str, force: bool, progress=None, fix_rotation: bool = True) -> tuple[int, int]:
    """Keep each original page unchanged and add only an invisible OCR text layer.
    Returns (OCR page count, copied page count)."""
    doc = pymupdf.open(src)
    if doc.needs_pass:
        doc.close()
        raise RuntimeError("The PDF is password-protected.")
    output = pymupdf.open()
    temp = dst.with_name(dst.stem + ".partial.pdf")
    ocr_count = copied_count = 0
    try:
        for position, page_number in enumerate(pages, 1):
            page = doc[page_number]
            output.insert_pdf(doc, from_page=page_number, to_page=page_number)
            has_text = len(page.get_text("text").strip()) >= MIN_TEXT_CHARS
            if has_text and not force:
                copied_count += 1
                status = "text layer found, copied"
            else:
                target = output[-1]
                rotation = target.rotation
                target.set_rotation(0)  # OCR in unrotated page coordinates
                page_dpi = safe_dpi(target, dpi)
                ocr_bytes = _ocr_bytes(target, 0, page_dpi, lang, tessdata)
                angle, display = 0, rotation
                if fix_rotation:
                    with pymupdf.open("pdf", ocr_bytes) as first:
                        detected = detect_reading_angle(target, first[0].get_text(), lang, tessdata)
                    if detected is not None:  # confident: show the text upright
                        angle = display = detected
                    if angle:
                        ocr_bytes = _ocr_bytes(target, angle, page_dpi, lang, tessdata)
                with pymupdf.open("pdf", ocr_bytes) as ocr_doc:
                    ocr_page = ocr_doc[0]
                    strip_images(ocr_page)  # discard the large raster, keep only the text
                    # The text layer was read from a turned image; turn it back onto the page.
                    target.show_pdf_page(target.rect, ocr_doc, 0, overlay=True, rotate=-angle)
                # Show the text upright when its direction is clear; otherwise keep the page's own turn.
                target.set_rotation(display)
                ocr_count += 1
                status = "OCR done" if page_dpi == dpi else f"OCR done (large page, read at {page_dpi} dpi)"
                if display != rotation:
                    status += " (sideways page turned upright)"
            if progress:
                progress(position, len(pages), page_number + 1, status)
        output.save(temp, garbage=4, deflate=True, use_objstms=True)
        output.close()
        os.replace(temp, dst)  # the final file only appears once it is complete
    finally:
        if not output.is_closed:
            output.close()
        doc.close()
        if temp.exists():
            try:
                temp.unlink()
            except OSError:
                pass
    return ocr_count, copied_count


# =============================================================== document library
_LIBRARY_LETTERS = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ى": "ي", "ة": "ه", "ؤ": "و", "ئ": "ي"})


def normalize_search(text: str) -> str:
    """Search form of a text: lower case, Western digits, no tatweel/diacritics, unified Arabic letters."""
    import unicodedata
    text = unicodedata.normalize("NFKC", text)  # ligatures such as "ﬂ" and Arabic presentation forms
    text = text.translate(ARABIC_DIGITS).replace(TATWEEL, "")
    text = ARABIC_DIACRITICS.sub("", text).translate(_LIBRARY_LETTERS).lower()
    return re.sub(r"\s+", " ", text)


def page_texts(pdf_path: Path) -> list[str]:
    """Readable text of each page; OCR text layers get Arabic words back in reading order."""
    texts = []
    with pymupdf.open(pdf_path) as doc:
        for page in doc:
            text = page.get_text("text", sort=True)
            if is_ocr_page(page):
                text = "\n".join(fix_rtl_line(line) for line in text.splitlines())
            texts.append(text)
    return texts


class DocumentLibrary:
    """Offline index of processed documents, stored in one SQLite file next to the settings."""

    def __init__(self, path: Path | None = None) -> None:
        self.path = path or settings_path().with_name("library.db")
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS documents (
                    id INTEGER PRIMARY KEY, path TEXT UNIQUE, title TEXT, pages INTEGER,
                    added REAL, modified REAL, size INTEGER);
                CREATE TABLE IF NOT EXISTS pages (doc_id INTEGER, page INTEGER, text TEXT, norm TEXT);
                CREATE INDEX IF NOT EXISTS pages_doc ON pages(doc_id);
            """)

    def _connect(self):
        import sqlite3
        db = sqlite3.connect(self.path, timeout=10)
        db.execute("PRAGMA foreign_keys = ON")
        return db

    def add(self, path: Path, texts: list[str]) -> None:
        """Add or refresh a document. `path` is the file the user should open from a search result."""
        path = Path(path).resolve()
        stat = path.stat()
        with self._connect() as db:
            old = db.execute("SELECT id FROM documents WHERE path = ?", (str(path),)).fetchone()
            if old:
                db.execute("DELETE FROM pages WHERE doc_id = ?", old)
                db.execute("DELETE FROM documents WHERE id = ?", old)
            cursor = db.execute(
                "INSERT INTO documents (path, title, pages, added, modified, size) VALUES (?, ?, ?, ?, ?, ?)",
                (str(path), path.name, len(texts), time.time(), stat.st_mtime, stat.st_size))
            db.executemany("INSERT INTO pages (doc_id, page, text, norm) VALUES (?, ?, ?, ?)",
                           [(cursor.lastrowid, number, text, normalize_search(text))
                            for number, text in enumerate(texts, 1)])

    def contains(self, path: Path) -> bool:
        with self._connect() as db:
            return db.execute("SELECT 1 FROM documents WHERE path = ?",
                              (str(Path(path).resolve()),)).fetchone() is not None

    def search(self, query: str, limit: int = 500) -> list[dict]:
        """All words must appear on the same page; "quoted text" must appear as written."""
        phrases = [normalize_search(p) for p in re.findall(r'"([^"]+)"', query)]
        words = [normalize_search(w) for w in re.sub(r'"[^"]*"', " ", query).split()]
        terms = [t.strip() for t in phrases + words if t.strip()]
        if not terms:
            return []
        clause = " AND ".join("p.norm LIKE ? ESCAPE '\\'" for _ in terms)
        values = ["%" + t.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%" for t in terms]
        with self._connect() as db:
            rows = db.execute(
                f"SELECT d.id, d.path, d.title, p.page, p.norm FROM pages p JOIN documents d ON d.id = p.doc_id "
                f"WHERE {clause} ORDER BY d.added DESC, p.page LIMIT ?", (*values, limit)).fetchall()
        results = []
        for doc_id, path, title, page, norm in rows:
            at = norm.find(terms[0])
            start = max(0, at - 20)
            while start > 0 and norm[start - 1] not in " \n":  # start at a word boundary
                start -= 1
            snippet = ("…" if start else "") + norm[start:at + len(terms[0]) + 140].strip() + "…"
            results.append({"doc_id": doc_id, "path": path, "title": title, "page": page,
                            "snippet": snippet.replace("\n", " "), "exists": Path(path).exists()})
        return results

    def documents(self) -> list[dict]:
        with self._connect() as db:
            rows = db.execute("SELECT id, path, title, pages, added FROM documents ORDER BY added DESC").fetchall()
        return [{"id": r[0], "path": r[1], "title": r[2], "pages": r[3], "added": r[4],
                 "exists": Path(r[1]).exists()} for r in rows]

    def remove(self, doc_ids: list[int]) -> int:
        with self._connect() as db:
            for doc_id in doc_ids:
                db.execute("DELETE FROM pages WHERE doc_id = ?", (doc_id,))
                db.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
        return len(doc_ids)

    def remove_missing(self) -> int:
        return self.remove([d["id"] for d in self.documents() if not d["exists"]])

    def stats(self) -> tuple[int, int]:
        with self._connect() as db:
            return (db.execute("SELECT COUNT(*) FROM documents").fetchone()[0],
                    db.execute("SELECT COUNT(*) FROM pages").fetchone()[0])


# =============================================================== tables to Excel
def _runs(mask_line, min_len: int):
    """Start/end indexes of runs of True at least min_len long."""
    import numpy as np
    padded = np.concatenate(([0], mask_line.astype(np.int8), [0]))
    diff = np.diff(padded)
    starts, ends = np.where(diff == 1)[0], np.where(diff == -1)[0]
    keep = (ends - starts) >= min_len
    return starts[keep], ends[keep]


def _merge_segments(segments, tolerance: int = 4):
    """Merge parallel pixel runs (a thick line gives several) into single lines (pos, start, end)."""
    merged = []
    for pos, start, end in sorted(segments):
        for line in merged:
            if abs(line[0] - pos) <= tolerance and start <= line[2] + tolerance and end >= line[1] - tolerance:
                line[0] = (line[0] + pos) / 2
                line[1], line[2] = min(line[1], start), max(line[2], end)
                break
        else:
            merged.append([pos, start, end])
    return merged


def _cluster(values, tolerance: float):
    out = []
    for value in sorted(values):
        if out and value - out[-1][-1] <= tolerance:
            out[-1].append(value)
        else:
            out.append([value])
    return [sum(group) / len(group) for group in out]


def _scan_tables(page, lang: str, tessdata: str, dpi: int = 200) -> list[list[list[str]]]:
    """Read ruled tables from a scanned page: find the grid lines in the image, erase them,
    OCR the clean image once, and place each word into its grid cell."""
    import numpy as np
    zoom = dpi / 72
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), colorspace=pymupdf.csGRAY)
    image = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.stride)[:, :pix.width].copy()
    dark = image < 150
    height, width = dark.shape
    min_h, min_v = max(40, int(width * 0.04)), max(25, int(height * 0.012))
    h_runs = [(y, s, e) for y in range(height) for s, e in zip(*_runs(dark[y], min_h))]
    v_runs = [(x, s, e) for x in range(width) for s, e in zip(*_runs(dark[:, x], min_v))]
    if not h_runs or not v_runs:
        return []
    erase = np.zeros_like(dark)
    for y, s, e in h_runs:
        erase[max(0, y - 1):y + 2, s:e] = True
    for x, s, e in v_runs:
        erase[s:e, max(0, x - 1):x + 2] = True
    h_lines, v_lines = _merge_segments(h_runs), _merge_segments(v_runs)

    # Group lines that touch each other into separate tables.
    tol = 6
    nodes = [("h", i) for i in range(len(h_lines))] + [("v", i) for i in range(len(v_lines))]
    parent = {node: node for node in nodes}

    def find(node):
        while parent[node] != node:
            parent[node] = parent[parent[node]]
            node = parent[node]
        return node
    for i, (hy, hx0, hx1) in enumerate(h_lines):
        for j, (vx, vy0, vy1) in enumerate(v_lines):
            if hx0 - tol <= vx <= hx1 + tol and vy0 - tol <= hy <= vy1 + tol:
                parent[find(("h", i))] = find(("v", j))
    groups: dict = {}
    for node in nodes:
        groups.setdefault(find(node), []).append(node)

    candidates = []
    for members in groups.values():
        hs = [h_lines[i] for kind, i in members if kind == "h"]
        vs = [v_lines[i] for kind, i in members if kind == "v"]
        ys, xs = _cluster([h[0] for h in hs], tol), _cluster([v[0] for v in vs], tol)
        if len(ys) >= 2 and len(xs) >= 2:
            candidates.append((ys, xs, vs))
    if not candidates:
        return []

    image[erase] = 255  # remove the grid so OCR does not read lines as characters

    def read_words(region, offset_x: int = 0, offset_y: int = 0, language: str = lang):
        """OCR a grayscale region; returns words as (x0, x1, centre x, centre y, height, text) in image pixels."""
        padded = np.full((region.shape[0] + 40, region.shape[1] + 40), 255, np.uint8)
        padded[20:-20, 20:-20] = region
        # The OCR engine reads colour (RGB) images only.
        picture = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.Pixmap(pymupdf.csGRAY, padded.shape[1], padded.shape[0],
                                                                padded.tobytes(), 0))
        picture.set_dpi(dpi, dpi)
        with pymupdf.open("pdf", picture.pdfocr_tobytes(language=language, tessdata=tessdata)) as ocr:
            scale = padded.shape[1] / ocr[0].rect.width  # OCR page points -> pixels
            out = []
            for w in ocr[0].get_text("words"):
                x0, x1 = w[0] * scale - 20 + offset_x, w[2] * scale - 20 + offset_x
                y0, y1 = w[1] * scale - 20 + offset_y, w[3] * scale - 20 + offset_y
                out.append((x0, x1, (x0 + x1) / 2, (y0 + y1) / 2, y1 - y0, w[4]))
            return out

    def cell_text(found) -> str:
        # left-to-right within each text line; _clean_cell restores Arabic reading order
        ordered = sorted(found, key=lambda w: (round(w[3] / max(w[4] * 0.8, 1)), w[2]))
        return _clean_cell(" ".join(w[5] for w in ordered))

    words = read_words(image)
    # The Arabic model can misread Western digits; an English-only pass gives a second opinion for numbers.
    digit_words = read_words(image, language="eng") if "ara" in lang.split("+") and "eng" in lang.split("+") else []

    tables = []
    for ys, xs, vs in sorted(candidates, key=lambda c: c[0][0]):
        rows = []
        for top, bottom in zip(ys, ys[1:]):
            middle = (top + bottom) / 2
            # A column boundary only counts in this row if a vertical line crosses the row.
            boundaries = [xs[0]] + [x for x in xs[1:-1]
                                    if any(abs(v[0] - x) <= tol and v[1] - tol <= middle <= v[2] + tol for v in vs)] + [xs[-1]]
            def place(found):
                cells, crossing = [[] for _ in xs[1:]], set()
                for word in found:
                    x0, x1, cx, cy = word[:4]
                    if top <= cy <= bottom and xs[0] <= cx <= xs[-1]:
                        # the (merged) cell that contains the word, under the first grid column it spans
                        for b_left, b_right in zip(boundaries, boundaries[1:]):
                            if b_left <= cx <= b_right:
                                column = xs.index(b_left)
                                cells[column].append(word)
                                if x0 < b_left - tol or x1 > b_right + tol:
                                    crossing.add(column)  # OCR joined text from neighbouring cells
                                break
                return cells, crossing

            cells, spans = place(words)
            row = [cell_text(cell) for cell in cells]
            if digit_words:
                digit_cells, digit_spans = place(digit_words)
                for column, found in enumerate(digit_cells):
                    candidate = cell_text(found)
                    if column in digit_spans or any(column + k in digit_spans for k in (-1, 1)):
                        # the English pass joined neighbouring numbers: read this cell on its own
                        b_left = xs[column]
                        b_right = next((b for b in boundaries if b > b_left + tol), xs[-1])
                        region = image[int(top) + 3:int(bottom) - 2, int(b_left) + 3:int(b_right) - 2]
                        candidate = cell_text(read_words(region, int(b_left) + 3, int(top) + 3, "eng")) \
                            if region.size else ""
                        digit_spans.discard(column)
                    if (column not in digit_spans and _to_number(candidate) is not None
                            and _to_number(row[column]) != _to_number(candidate)
                            and not re.search(r"[\u0600-\u06FF]{2,}", row[column])):
                        row[column] = candidate
            if spans:
                # Re-read every cell of this row on its own, so text cannot run across column lines.
                for b_left, b_right in zip(boundaries, boundaries[1:]):
                    column = xs.index(b_left)
                    region = image[int(top) + 3:int(bottom) - 2, int(b_left) + 3:int(b_right) - 2]
                    if region.size:
                        row[column] = cell_text(read_words(region, int(b_left) + 3, int(top) + 3))
            if any(row):
                rows.append(row)
        if len(rows) >= 2:
            tables.append(rows)
    return tables


def _to_number(text: str):
    """Return a number for cells that are plain amounts (e.g. "1,250.00", "(45)", "١٢٠"), else None."""
    value = text.translate(ARABIC_DIGITS).strip().strip("|[]_ ").replace(" ", "")
    negative = value.startswith("(") and value.endswith(")")
    value = value.strip("()")
    if not re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?|-?\d+(\.\d+)?", value):
        return None
    number = float(value.replace(",", ""))
    if negative:
        number = -number
    return int(number) if number.is_integer() and "." not in value else number


def _clean_cell(text: str) -> str:
    text = re.sub(r"\s+", " ", text or "").strip().strip("|[]_").strip()
    return fix_rtl_line(text) if ARABIC.search(text) else text


def extract_tables(pdf_path: Path, page_numbers: list[int], lang: str, tessdata: str, progress=None) -> list[dict]:
    """Find ruled tables. Digital pages use the drawn lines; scanned pages use lines found in the image,
    which are erased before the cell text is read. Returns [{page, index, rows}]."""
    found = []
    with pymupdf.open(pdf_path) as doc:
        for position, index in enumerate(page_numbers):
            page = doc[index]
            scanned = is_ocr_page(page) or len(page.get_text("text").strip()) < MIN_TEXT_CHARS
            tables = []
            if not scanned:
                for table in page.find_tables(strategy="lines").tables:
                    tables.append([[_clean_cell(cell) for cell in row] for row in table.extract()])
            else:
                tables = _scan_tables(page, lang, tessdata)
            for number, rows in enumerate(tables, 1):
                if any(any(cell for cell in row) for row in rows):
                    found.append({"page": page_numbers[position] + 1, "index": number, "rows": rows})
            if progress:
                progress(position + 1, len(page_numbers), page_numbers[position] + 1,
                         f"{len(tables)} table(s) found" if tables else "no ruled table found")
    return found


def write_tables_xlsx(tables: list[dict], path: Path, title: str) -> None:
    """One sheet per table (named P<page>-T<number>), numbers stored as numbers, plus an index sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    index = workbook.active
    index.title = "Tables"
    index.append([f"{title}: tables found by GrepAll"])
    index["A1"].font = Font(bold=True, size=12)
    index.append([])
    index.append(["Sheet", "Page", "Rows", "Columns"])
    for cell in index[3]:
        cell.font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    for table in tables:
        name = f"P{table['page']}-T{table['index']}"
        sheet = workbook.create_sheet(name)
        rows = table["rows"]
        widths: dict[int, int] = {}
        arabic_cells = 0
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row, 1):
                number = _to_number(value) if r > 1 else None
                cell = sheet.cell(row=r, column=c, value=number if number is not None else (value or None))
                if number is not None and isinstance(number, float):
                    cell.number_format = "#,##0.00"
                elif number is not None:
                    cell.number_format = "#,##0"
                if r == 1:
                    cell.font = Font(bold=True)
                    cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                widths[c] = max(widths.get(c, 0), min(60, len(str(value or ""))))
                arabic_cells += 1 if ARABIC.search(value or "") else 0
        for c, width in widths.items():
            sheet.column_dimensions[get_column_letter(c)].width = max(8, width + 2)
        sheet.freeze_panes = "A2"
        if arabic_cells > sum(len(row) for row in rows) / 3:
            sheet.sheet_view.rightToLeft = True
        index.append([name, table["page"], len(rows), max((len(row) for row in rows), default=0)])
    index.column_dimensions["A"].width = 14
    temp = path.with_name(path.stem + ".partial.xlsx")
    workbook.save(temp)
    os.replace(temp, path)


# =============================================================== reduce file size
REDUCE_LEVELS = {
    # name: (description, image settings or None for structure-only clean-up)
    "Lossless": ("File structure only. Quality unchanged; savings are usually small.", None),
    "Balanced": ("Images at 200 dpi, high quality. Text stays sharp. Recommended.",
                 dict(dpi_threshold=225, dpi_target=200, quality=80)),
    "Smallest": ("Images at 150 dpi, stronger compression, e.g. for email.",
                 dict(dpi_threshold=165, dpi_target=150, quality=65)),
}


def reduce_pdf(src: Path, dst: Path, level: str = "Balanced", grayscale: bool = False) -> tuple[int, int, bool]:
    """Write a smaller copy of a PDF. Returns (size before, size after, saved).
    Nothing is saved when the result would not be smaller than the original."""
    before = src.stat().st_size
    doc = pymupdf.open(src)
    if doc.needs_pass:
        doc.close()
        raise RuntimeError("The PDF is password-protected.")
    temp = dst.with_name(dst.stem + ".partial.pdf")
    try:
        settings = REDUCE_LEVELS[level][1]
        if settings or grayscale:
            # Black-and-white scans are already stored very compactly, so they are left as they are.
            doc.rewrite_images(**(settings or dict(dpi_threshold=10000, dpi_target=0, quality=90)),
                               bitonal=False, set_to_gray=grayscale)
        try:
            doc.subset_fonts()  # keep only the characters that are used
        except Exception:
            pass
        doc.save(temp, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True,
                 use_objstms=True, clean=True)
        doc.close()
        after = temp.stat().st_size
        if after >= before:
            temp.unlink()
            return before, before, False
        os.replace(temp, dst)
        return before, after, True
    finally:
        if not doc.is_closed:
            doc.close()
        if temp.exists():
            try:
                temp.unlink()
            except OSError:
                pass


def human_size(size: int) -> str:
    for unit in ("bytes", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{size:.0f} {unit}" if unit == "bytes" else f"{size:.1f} {unit}"
        size /= 1024
    return f"{size:.1f} GB"


def is_ocr_page(page) -> bool:
    """True when the page's text layer was produced by Tesseract (invisible GlyphLessFont)."""
    return any("GlyphLess" in font[3] for font in page.get_fonts(full=True))


def fix_rtl_line(line: str) -> str:
    """Tesseract stores Arabic words in visual order; restore logical order for
    Arabic-dominant lines while keeping English words and number sequences intact."""
    prefix, body = MD_PREFIX.match(line).groups()
    if len(ARABIC.findall(body)) <= len(re.findall(r"[A-Za-z]", body)):
        return line
    groups: list[list[str]] = []
    for token in body.split():
        ltr_token = bool(LTR_CHARS.search(token) and not ARABIC.search(token))
        if ltr_token and groups and groups[-1][0] == "L":
            groups[-1].append(token)
        else:
            groups.append(["L" if ltr_token else "R", token])
    return prefix + " ".join(" ".join(group[1:]) for group in reversed(groups))


# =============================================================== sensitive data
# Order matters: longer, more specific patterns first, so their digits are not matched again.
SENSITIVE_PATTERNS = (
    ("IBAN", re.compile(r"\b[A-Z5][A-Z48]\d{2}(?:[ ]?[A-Z0-9]){11,30}\b")),
    ("CARD", re.compile(r"(?<!\d)\d(?:[ -]?\d){12,18}(?!\d)")),
    ("EMAIL", re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")),
    ("PHONE", re.compile(r"(?<![\w+])(?:(?:\+|00)?966[ -]?|0)5\d(?:[ -]?\d){7}(?!\d)"
                         r"|(?<!\w)\+\d{1,3}(?:[ -]?\d){6,12}(?!\d)")),
    ("ID", re.compile(r"(?<!\d)[12]\d{9}(?!\d)")),  # Saudi national ID / Iqama
)
REDACTION_NOTE = ("Automatic detection can miss items, especially on poor scans. "
                  "Review redacted files before sharing them.")


def _iban_valid(text: str) -> bool:
    value = re.sub(r"\s", "", text).upper()
    if re.fullmatch(r"[S5][A48]\d{22}", value):
        return True  # Saudi IBAN, including common OCR misreads of "SA" (5A, S8, 58...)
    if not value[:2].isalpha():
        return False
    if not 15 <= len(value) <= 34:
        return False
    digits = "".join(str(int(ch, 36)) for ch in value[4:] + value[:4])
    return int(digits) % 97 == 1


def _luhn_valid(text: str) -> bool:
    digits = [int(ch) for ch in text if ch.isdigit()]
    if not 13 <= len(digits) <= 19:
        return False
    total = 0
    for index, digit in enumerate(reversed(digits)):
        if index % 2:
            digit = digit * 2 - 9 if digit > 4 else digit * 2
        total += digit
    return total % 10 == 0


VALIDATORS = {"IBAN": _iban_valid, "CARD": _luhn_valid}


def parse_terms(raw: str) -> list[str]:
    """Custom words to hide, separated by commas or new lines."""
    return [term.strip() for term in re.split(r"[,\n،]", raw) if len(term.strip()) >= 2]


_MONTHS = ("January|February|March|April|May|June|July|August|September|October|November|December|"
           "Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec|"
           "يناير|فبراير|مارس|أبريل|ابريل|مايو|يونيو|يوليو|أغسطس|اغسطس|سبتمبر|أكتوبر|اكتوبر|نوفمبر|ديسمبر")
_WEEKDAY = (r"(?:(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday|Mon|Tue|Tues|Wed|Thu|Thur|Thurs|Fri|Sat|Sun|"
            r"الأحد|الاحد|الاثنين|الإثنين|الثلاثاء|الأربعاء|الاربعاء|الخميس|الجمعة|السبت|يوم)\.?,?\s+)?")
# A value is a date (numeric or with a month name) or, otherwise, one "word" containing a digit,
# such as HMG-DAM-MEP-TEN-01. The digit rule keeps ordinary words ("dated in the past") visible.
_VALUE = (rf"(?:{_WEEKDAY}\d{{1,4}}[/.\-]\d{{1,2}}[/.\-]\d{{1,4}}(?:\s*(?:هـ|ه|م|AH|H))?"
          rf"|{_WEEKDAY}\d{{1,2}}(?:st|nd|rd|th)?\s+(?:{_MONTHS})\.?,?\s+\d{{2,4}}"
          rf"|{_WEEKDAY}(?:{_MONTHS})\.?\s+\d{{1,2}}(?:st|nd|rd|th)?,?\s+\d{{4}}"
          r"|[^\s,;:()\[\]]*\d(?:[^\s,;:()\[\]]*[^\s,;:.()\[\]])?)")


def parse_labels(raw: str) -> list[str]:
    """Words after which the value is hidden, e.g. "Ref, Dated, Contract No"."""
    return parse_terms(raw)


def _label_matcher(labels: list[str]):
    """Match a label and capture the value that follows it (group "v"); the label itself stays visible."""
    names = [r"\s*".join(re.escape(part) for part in label.split())
             for label in sorted(labels, key=len, reverse=True)]
    any_label = "|".join(names)
    return re.compile(
        rf"(?<!\w)(?:{any_label})(?!\w)\.?"                     # the label, e.g. "Ref" or "Ref."
        r"(?:\s*(?:No|Number|Nr|رقم)\.?)?\s*[:#\-–]?\s*"          # optional "No." and separator
        rf"(?!(?:{any_label})(?!\w))"                             # the next word is not another label
        rf"(?P<v>{_VALUE})", re.IGNORECASE)


def _matchers(terms: list[str], labels: list[str] = ()):
    matchers = list(SENSITIVE_PATTERNS)
    if labels:
        matchers.insert(0, ("REDACTED", _label_matcher(labels)))
    if terms:
        alternatives = "|".join(r"\s*".join(re.escape(part) for part in term.split())
                                for term in sorted(terms, key=len, reverse=True))
        matchers.insert(0, ("REDACTED", re.compile(alternatives, re.IGNORECASE)))
    return matchers


def _find_sensitive(text: str, terms: list[str], labels: list[str] = ()):
    """Yield (label, start, end) for each item to hide, without overlaps."""
    taken: list[tuple[int, int]] = []
    found = []
    search = text.translate(ARABIC_DIGITS)  # one-to-one mapping, so positions stay the same
    for label, pattern in _matchers(terms, labels):
        validate = VALIDATORS.get(label)
        for match in pattern.finditer(search):
            start, end = match.span("v") if "v" in pattern.groupindex else match.span()
            if validate and not validate(match.group()):
                continue
            if any(start < other_end and end > other_start for other_start, other_end in taken):
                continue
            taken.append((start, end))
            found.append((label, start, end))
    return sorted(found, key=lambda item: item[1])


def redact_text(text: str, terms: list[str], labels: list[str] = ()) -> tuple[str, Counter]:
    counts: Counter = Counter()
    parts, position = [], 0
    for label, start, end in _find_sensitive(text, terms, labels):
        parts.append(text[position:start])
        parts.append(f"[{label}]")
        counts[label] += 1
        position = end
    parts.append(text[position:])
    result = "".join(parts)
    # Safety net for multi-word custom terms: OCR can join or reorder words in mixed Arabic lines.
    # When every word of a term appears on the same line, hide each of them.
    for term in terms:
        words = [word for word in term.split() if len(word) >= 2]
        if len(words) < 2:
            continue
        pattern = re.compile("|".join(re.escape(word) for word in words), re.IGNORECASE)
        lines = result.split("\n")
        for index, line in enumerate(lines):
            lowered = line.lower()
            if all(word.lower() in lowered for word in words):
                lines[index], count = pattern.subn("[REDACTED]", line)
                counts["REDACTED"] += 1 if count else 0
        result = "\n".join(lines)
    return result, counts


def redact_pdf(src: Path, dst: Path, terms: list[str], labels: list[str] = ()) -> int:
    """Create a copy with sensitive items permanently removed (text and image pixels).
    Returns the number of areas blacked out."""
    doc = pymupdf.open(src)
    areas = 0
    try:
        for page in doc:
            words = page.get_text("words")
            lines: dict = {}
            for word in words:
                lines.setdefault((word[5], word[6]), []).append(word)
            selected: set[int] = set()
            index_of = {id(word): index for index, word in enumerate(words)}
            ocr_page = is_ocr_page(page)
            for line in lines.values():
                line.sort(key=lambda word: word[7])
                # OCR stores Arabic words in visual order; also try the reversed order.
                for sequence in ([line, list(reversed(line))] if ocr_page else [line]):
                    text, spans = "", []
                    for word in sequence:
                        spans.append((len(text), len(text) + len(word[4])))
                        text += word[4] + " "
                    for _label, start, end in _find_sensitive(text, terms, labels):
                        for word, (word_start, word_end) in zip(sequence, spans):
                            if word_start < end and word_end > start:
                                selected.add(index_of[id(word)])
            if selected:
                strip_images(page, tiny_only=True)  # placeholders from v1.4 and earlier OCR files
            for index in selected:
                rect = pymupdf.Rect(words[index][:4])
                page.add_redact_annot(rect + (-2, -2, 2, 2), fill=(0, 0, 0))
            if selected:
                areas += len(selected)
                page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS)
        temp = dst.with_name(dst.stem + ".partial.pdf")
        doc.save(temp, garbage=4, deflate=True)
        doc.close()
        os.replace(temp, dst)
    finally:
        if not doc.is_closed:
            doc.close()
    return areas


def to_markdown(pdf_path: Path, md_path: Path, title: str, page_numbers: list[int],
                markers: bool, strip_headers: bool = False, cleanup: dict | None = None,
                redact_terms: list[str] | None = None,
                redact_labels: list[str] = ()) -> tuple[str | None, Counter]:
    """Write Markdown. redact_terms (a list, possibly empty) turns on sensitive data hiding.
    Returns (warning when the layout engine was not used, counts of hidden items)."""
    counts: Counter = Counter()
    texts: list[str] = []
    warning = None
    try:
        import pymupdf4llm
        common = dict(page_chunks=True, write_images=False, show_progress=False)
        try:
            data = pymupdf4llm.to_markdown(str(pdf_path), use_ocr=False, header=not strip_headers,
                                           footer=not strip_headers, **common)
        except TypeError:
            data = pymupdf4llm.to_markdown(str(pdf_path), ignore_images=True, **common)
        texts = [chunk.get("text", "") for chunk in data]
    except ImportError:
        warning = "pymupdf4llm is not installed; plain text layout was used."
    except Exception as exc:
        warning = f"Layout conversion failed ({exc}); plain text layout was used."

    chunks: list[str] = []
    with pymupdf.open(pdf_path) as doc:
        for index, page in enumerate(doc):
            body = texts[index].strip() if index < len(texts) else ""
            if not body:
                body = page.get_text("text", sort=True).strip() or "_[no text on this page]_"
            if is_ocr_page(page):
                body = "\n".join(fix_rtl_line(line) for line in body.splitlines())
            if cleanup:
                body = clean_arabic(body, **cleanup)
            if redact_terms is not None:
                body, found = redact_text(body, redact_terms, redact_labels)
                counts.update(found)
            number = page_numbers[index] if index < len(page_numbers) else index + 1
            header = f"<!-- Page {number} -->\n\n" if markers else ""
            chunks.append(header + body)
    temp = md_path.with_name(md_path.name + ".partial")
    temp.write_text(f"# {title}\n\n" + "\n\n".join(chunks) + "\n", encoding="utf-8")
    os.replace(temp, md_path)
    return warning, counts


# =============================================================== Word, Excel and PowerPoint
OFFICE_KINDS = {".docx": "Word", ".xlsx": "Excel", ".xlsm": "Excel", ".pptx": "PowerPoint"}
MAX_SHEET_ROWS = 20_000  # per sheet; keeps the Markdown usable


def is_office(path: Path) -> bool:
    return path.suffix.lower() in OFFICE_KINDS


def _md_cell(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", str(text).replace("\r", ""))
    return "<br>".join(part.strip() for part in text.split("\n") if part.strip()).replace("|", "\\|")


def _md_table(rows: list[list[str]]) -> str:
    """Markdown table without empty rows and columns; the first row is the header."""
    rows = [[_md_cell(cell) for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    keep = [i for i in range(width) if any(row[i] for row in rows)]
    rows = [[row[i] for i in keep] for row in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * len(keep)]
    lines += ["| " + " | ".join(row) + " |" for row in rows[1:]]
    return "\n".join(lines)


def _styled(pieces: list[tuple[str, bool, bool]]) -> str:
    """Join runs, merging neighbours with the same bold/italic so markers stay clean."""
    merged: list[list] = []
    for text, bold, italic in pieces:
        if merged and merged[-1][1:] == [bold, italic]:
            merged[-1][0] += text
        else:
            merged.append([text, bold, italic])
    out = []
    for text, bold, italic in merged:
        core = text.strip()
        if not core or not (bold or italic):
            out.append(text)
            continue
        mark = "***" if bold and italic else "**" if bold else "*"
        lead, trail = text[:len(text) - len(text.lstrip())], text[len(text.rstrip()):]
        out.append(f"{lead}{mark}{core}{mark}{trail}")
    return "".join(out)


class _DocxNumbering:
    """Rebuilds Word's automatic numbering (1.2.3, (a), iv., bullets) so clause numbers survive."""

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def __init__(self, doc) -> None:
        self.levels: dict[str, dict[int, tuple[int, str, str]]] = {}  # numId -> level -> (start, fmt, text)
        self.counters: dict[str, dict[int, int]] = {}
        w = self.W
        try:
            root = doc.part.numbering_part.element
        except Exception:
            return
        abstracts = {}
        for abstract in root.iter(f"{w}abstractNum"):
            levels = {}
            for lvl in abstract.iter(f"{w}lvl"):
                ilvl = int(lvl.get(f"{w}ilvl", "0"))
                start = lvl.find(f"{w}start")
                fmt = lvl.find(f"{w}numFmt")
                text = lvl.find(f"{w}lvlText")
                levels[ilvl] = (int(start.get(f"{w}val")) if start is not None else 1,
                                fmt.get(f"{w}val") if fmt is not None else "decimal",
                                text.get(f"{w}val") if text is not None else f"%{ilvl + 1}.")
            abstracts[abstract.get(f"{w}abstractNumId")] = levels
        for num in root.iter(f"{w}num"):
            ref = num.find(f"{w}abstractNumId")
            levels = dict(abstracts.get(ref.get(f"{w}val") if ref is not None else "", {}))
            for override in num.iter(f"{w}lvlOverride"):
                ilvl = int(override.get(f"{w}ilvl", "0"))
                start = override.find(f"{w}startOverride")
                if start is not None and ilvl in levels:
                    levels[ilvl] = (int(start.get(f"{w}val")),) + levels[ilvl][1:]
            self.levels[num.get(f"{w}numId")] = levels

    def _numpr(self, paragraph) -> tuple[str | None, int | None]:
        num_id = level = None
        sources = [paragraph._p.pPr]
        style = paragraph.style
        while style is not None:
            sources.append(style.element.pPr)
            style = style.base_style
        for ppr in sources:
            numpr = ppr.numPr if ppr is not None else None
            if numpr is None:
                continue
            if num_id is None and numpr.numId is not None:
                num_id = str(numpr.numId.val)
            if level is None and numpr.ilvl is not None:
                level = numpr.ilvl.val
        return num_id, level

    @staticmethod
    def _format(value: int, fmt: str) -> str:
        if fmt in ("lowerLetter", "upperLetter"):
            text = ""
            while value > 0:
                value, rest = divmod(value - 1, 26)
                text = chr(97 + rest) + text
            return text.upper() if fmt == "upperLetter" else text
        if fmt in ("lowerRoman", "upperRoman"):
            pairs = ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"), (50, "l"),
                     (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"))
            text = ""
            for number, letters in pairs:
                while value >= number:
                    text, value = text + letters, value - number
            return text.upper() if fmt == "upperRoman" else text
        if fmt in ("arabicAlpha", "arabicAbjad"):
            letters = ("أبتثجحخدذرزسشصضطظعغفقكلمنهوي" if fmt == "arabicAlpha" else "أبجدهوزحطيكلمنسعفصقرشتثخذضظغ")
            return letters[(value - 1) % len(letters)]
        return str(value)

    def label(self, paragraph) -> tuple[int, str, bool] | None:
        """(level, label text, is bullet) for automatically numbered paragraphs, otherwise None."""
        num_id, level = self._numpr(paragraph)
        if not num_id or num_id == "0" or num_id not in self.levels:
            return None
        level = level or 0
        levels = self.levels[num_id]
        start, fmt, text = levels.get(level, (1, "decimal", f"%{level + 1}."))
        if fmt == "none":
            return None
        counters = self.counters.setdefault(num_id, {})
        counters[level] = counters.get(level, start - 1) + 1
        for deeper in [k for k in counters if k > level]:
            del counters[deeper]
        if fmt == "bullet":
            return level, "", True

        def substitute(match: re.Match) -> str:
            index = int(match.group(1)) - 1
            lvl_start, lvl_fmt, _ = levels.get(index, (1, "decimal", ""))
            return self._format(counters.get(index, lvl_start), lvl_fmt)
        return level, re.sub(r"%(\d)", substitute, text).strip(), False


def _docx_table(table) -> str:
    rows, previous = [], {}
    for row in table.rows:
        try:
            cells = row.cells
        except Exception:
            continue
        values, seen = [], set()
        for column, cell in enumerate(cells):
            tc = id(cell._tc)
            # merged cells are repeated by python-docx: keep the text once (first column / first row)
            if tc in seen or previous.get(column) == tc:
                values.append("")
            else:
                values.append(cell.text)
            seen.add(tc)
            previous[column] = tc
        rows.append(values)
    return _md_table(rows)


def docx_sections(path: Path) -> tuple[list[tuple[str, str]], str]:
    import docx
    from docx.table import Table
    from docx.text.hyperlink import Hyperlink
    doc = docx.Document(str(path))
    blocks: list[str] = []
    numbering = _DocxNumbering(doc)
    paragraphs = tables = 0
    for item in doc.iter_inner_content():
        if isinstance(item, Table):
            text = _docx_table(item)
            if text:
                blocks.append(text)
                tables += 1
            continue
        pieces = []
        for part in item.iter_inner_content():
            runs = part.runs if isinstance(part, Hyperlink) else [part]
            for run in runs:
                pieces.append((run.text, bool(run.bold), bool(run.italic)))
        text = _styled(pieces).strip()
        numbered = numbering.label(item)  # counted even for empty paragraphs, as Word does
        if not text:
            continue
        paragraphs += 1
        style = (item.style.name if item.style is not None else "") or ""
        heading = re.match(r"Heading (\d)", style)
        if style == "Title" or heading:
            level = 1 if style == "Title" else int(heading.group(1))
            plain = re.sub(r"\*+", "", text).replace("\n", " ").strip()
            if numbered and numbered[1]:
                plain = f"{numbered[1]} {plain}"
            blocks.append("#" * min(level + 1, 6) + " " + plain)
            continue
        if numbered is None and style.startswith(("List Bullet", "List Number")):
            numbered = (0, "1." if style.startswith("List Number") else "", not style.startswith("List Number"))
        if numbered:
            level, label, bullet = numbered
            line = text.replace("\n", " ")
            if bullet:
                blocks.append("  " * level + "- " + line)
            elif re.fullmatch(r"\d+\.", label) and level == 0:
                blocks.append(f"{label} {line}")
            else:
                blocks.append("  " * level + f"- {label} {line}")
            continue
        blocks.append(text.replace("\n", "  \n"))
    # list items stay together; other blocks are separated by a blank line
    out: list[str] = []
    for block in blocks:
        is_item = bool(re.match(r"\s*(-|\d+\.) ", block))
        if out and is_item and re.match(r"\s*(-|\d+\.) ", out[-1].splitlines()[-1]):
            out[-1] += "\n" + block
        else:
            out.append(block)
    summary = f"Word: {paragraphs} paragraph(s), {tables} table(s)"
    return [("", "\n\n".join(out))], summary


def _excel_value(cell) -> str:
    import datetime as dt
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        fmt = getattr(cell, "number_format", "General") or "General"
        decimals = None
        match = re.search(r"0\.(0+)", fmt)
        if match:
            decimals = len(match.group(1))
        elif re.search(r"(^|[^.])0(%|$|[^0.])", fmt) or fmt in ("0", "#,##0"):
            decimals = 0
        if "%" in fmt:
            return f"{value * 100:.{decimals or 0}f}%"
        if decimals is not None:
            return f"{value:.{decimals}f}"
        if float(value).is_integer():
            return str(int(value))
        return format(round(float(value), 10), "f").rstrip("0").rstrip(".")
    return str(value)


def xlsx_sections(path: Path) -> tuple[list[tuple[str, str]], str]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(str(path), read_only=True, data_only=False)  # to spot formulas never calculated
    sections, total_rows, hidden, uncalculated = [], 0, 0, 0
    try:
        for ws, wf in zip(wb.worksheets, formulas.worksheets):
            if getattr(ws, "sheet_state", "visible") != "visible":
                hidden += 1
                continue
            rows, truncated, blank_run = [], False, 0
            for row, formula_row in zip(ws.iter_rows(), wf.iter_rows()):
                values = [_excel_value(cell) for cell in row]
                for cell, formula_cell in zip(row, formula_row):
                    if cell.value is None and isinstance(formula_cell.value, str) and formula_cell.value.startswith("="):
                        uncalculated += 1
                if any(values):
                    rows.extend([[]] * blank_run)
                    blank_run = 0
                    rows.append(values)
                else:
                    blank_run += 1
                    if blank_run > 5000 and rows:
                        break  # formatted but empty rows to the end of the sheet
                if len(rows) >= MAX_SHEET_ROWS:
                    truncated = True
                    break
            table = _md_table(rows)
            used = sum(1 for row in rows if any(row))
            total_rows += used
            body = table or "_[empty sheet]_"
            if truncated:
                body += f"\n\n_[only the first {MAX_SHEET_ROWS:,} rows are included]_"
            sections.append((f"Sheet: {ws.title}", body))
    finally:
        wb.close()
        formulas.close()
    summary = f"Excel: {len(sections)} sheet(s), {total_rows:,} row(s)"
    if hidden:
        summary += f", {hidden} hidden sheet(s) skipped"
    if uncalculated:
        summary += (f"; {uncalculated} formula result(s) missing (the file was not saved by Excel: "
                    "open it in Excel, save it, and convert again)")
    return sections, summary


def pptx_sections(path: Path) -> tuple[list[tuple[str, str]], str]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(path))
    sections = []

    def walk(shapes):
        ordered = sorted(shapes, key=lambda s: ((s.top or 0) // 20000, s.left or 0))
        for shape in ordered:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes)
            else:
                yield shape

    for number, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip().replace("\n", " ") if title_shape is not None else ""
        parts = []
        for shape in walk(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                table = _md_table([[cell.text for cell in row.cells] for row in shape.table.rows])
                if table:
                    parts.append(table)
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                chart = shape.chart
                name = chart.chart_title.text_frame.text.strip() if chart.has_title else ""
                parts.append(f"_[chart{': ' + name if name else ''}]_")
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                lines = []
                for paragraph in shape.text_frame.paragraphs:
                    text = _styled([(run.text, bool(run.font.bold), bool(run.font.italic))
                                    for run in paragraph.runs]).strip()
                    if text:
                        lines.append((paragraph.level, text))
                if len(lines) == 1:
                    parts.append(lines[0][1])
                elif lines:
                    parts.append("\n".join("   " * level + "- " + text for level, text in lines))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append("> **Notes:** " + notes.replace("\n", "\n> "))
        heading = f"Slide {number}" + (f": {title}" if title else "")
        sections.append((heading, "\n\n".join(parts) or "_[no text on this slide]_"))
    return sections, f"PowerPoint: {len(sections)} slide(s)"


def office_sections(path: Path) -> tuple[list[tuple[str, str]], str]:
    """(heading, Markdown body) per sheet/slide (one section for Word), and a short summary."""
    kind = OFFICE_KINDS.get(path.suffix.lower())
    if kind == "Word":
        return docx_sections(path)
    if kind == "Excel":
        return xlsx_sections(path)
    if kind == "PowerPoint":
        return pptx_sections(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def office_texts(path: Path) -> list[str]:
    """Searchable text per sheet/slide for the library."""
    sections, _ = office_sections(path)
    return [(heading + "\n" + body).strip() for heading, body in sections]


def office_to_markdown(path: Path, md_path: Path, title: str, cleanup: dict | None = None,
                       redact_terms: list[str] | None = None,
                       redact_labels: list[str] = ()) -> tuple[list[str], Counter, str]:
    """Write Markdown for a Word, Excel or PowerPoint file.
    Returns (unredacted texts for the library, counts of hidden items, summary)."""
    sections, summary = office_sections(path)
    counts: Counter = Counter()
    chunks, library = [], []
    for heading, body in sections:
        text = (f"## {heading}\n\n" if heading else "") + body
        library.append((heading + "\n" + body).strip())
        if cleanup:
            text = clean_arabic(text, **cleanup)
        if redact_terms is not None:
            text, found = redact_text(text, redact_terms, redact_labels)
            counts.update(found)
        chunks.append(text)
    temp = md_path.with_name(md_path.name + ".partial")
    temp.write_text(f"# {title}\n\n" + "\n\n".join(chunks).strip() + "\n", encoding="utf-8")
    os.replace(temp, md_path)
    return library, counts, summary


# =============================================================== feedback form
class FeedbackDialog(tk.Toplevel):
    """Collects a bug report or improvement suggestion and opens it in the user's email app."""

    TYPES = ("Bug report", "Improvement suggestion")

    def __init__(self, parent: "PdfOcrApp") -> None:
        super().__init__(parent)
        self.parent = parent
        self.title("Report a bug or suggest an improvement")
        self.transient(parent)
        self.resizable(True, True)
        self.minsize(520, 520)

        self.kind = tk.StringVar(value=self.TYPES[0])
        self.summary = tk.StringVar()
        self.name = tk.StringVar()
        self.include_system = tk.BooleanVar(value=True)
        self.include_log = tk.BooleanVar(value=True)

        body = ttk.Frame(self, padding=16)
        body.pack(fill="both", expand=True)
        body.columnconfigure(1, weight=1)
        body.rowconfigure(3, weight=1)
        body.rowconfigure(4, weight=1)

        ttk.Label(body, text="Type").grid(row=0, column=0, sticky="w", pady=4)
        types = ttk.Frame(body)
        types.grid(row=0, column=1, sticky="w", pady=4)
        for value in self.TYPES:
            ttk.Radiobutton(types, text=value, value=value, variable=self.kind,
                            command=self._toggle_steps).pack(side="left", padx=(0, 16))

        ttk.Label(body, text="Summary *").grid(row=1, column=0, sticky="w", pady=4)
        self.summary_entry = ttk.Entry(body, textvariable=self.summary)
        self.summary_entry.grid(row=1, column=1, sticky="ew", pady=4)

        ttk.Label(body, text="Your name").grid(row=2, column=0, sticky="w", pady=4)
        ttk.Entry(body, textvariable=self.name).grid(row=2, column=1, sticky="ew", pady=4)

        ttk.Label(body, text="Details *").grid(row=3, column=0, sticky="nw", pady=4)
        self.details = tk.Text(body, height=6, wrap="word")
        self.details.grid(row=3, column=1, sticky="nsew", pady=4)

        self.steps_label = ttk.Label(body, text="Steps to\nreproduce")
        self.steps_label.grid(row=4, column=0, sticky="nw", pady=4)
        self.steps = tk.Text(body, height=4, wrap="word")
        self.steps.grid(row=4, column=1, sticky="nsew", pady=4)
        self.steps.insert("1.0", "1. \n2. \n3. ")

        ttk.Checkbutton(body, text="Include app and system information", variable=self.include_system)\
            .grid(row=5, column=1, sticky="w", pady=(8, 0))
        ttk.Checkbutton(body, text="Include recent activity log", variable=self.include_log)\
            .grid(row=6, column=1, sticky="w")

        ttk.Label(body, text=f"The report opens as a new email to {FEEDBACK_EMAIL} in your email app. "
                             "Review it there, then select Send.",
                  style="Muted.TLabel", wraplength=460, justify="left")\
            .grid(row=7, column=0, columnspan=2, sticky="w", pady=(12, 0))

        buttons = ttk.Frame(body)
        buttons.grid(row=8, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(buttons, text="Cancel", command=self.destroy).pack(side="right")
        ttk.Button(buttons, text="Copy report", command=self.copy_report).pack(side="right", padx=8)
        ttk.Button(buttons, text="Open in email app", command=self.send).pack(side="right")

        for text_box in (self.details, self.steps):
            getattr(parent, "style_text", lambda _w: None)(text_box)
        palette = getattr(parent, "palette", None)
        if palette:
            self.configure(bg=palette["window"])
        self.bind("<Escape>", lambda _event: self.destroy())
        self.summary_entry.focus_set()
        self.grab_set()
        self._center_on_parent()

    def _center_on_parent(self) -> None:
        self.update_idletasks()
        x = self.parent.winfo_rootx() + (self.parent.winfo_width() - self.winfo_width()) // 2
        y = self.parent.winfo_rooty() + (self.parent.winfo_height() - self.winfo_height()) // 3
        self.geometry(f"+{max(x, 0)}+{max(y, 0)}")

    def _toggle_steps(self) -> None:
        state = "normal" if self.kind.get() == self.TYPES[0] else "disabled"
        self.steps.configure(state=state)

    def _validate(self) -> bool:
        if not self.summary.get().strip():
            messagebox.showwarning("Summary required", "Enter a short summary.", parent=self)
            self.summary_entry.focus_set()
            return False
        if not self.details.get("1.0", "end").strip():
            messagebox.showwarning("Details required", "Describe the issue or suggestion.", parent=self)
            self.details.focus_set()
            return False
        return True

    def _subject(self) -> str:
        tag = "Bug" if self.kind.get() == self.TYPES[0] else "Suggestion"
        return f"[{APP_NAME} {APP_VERSION}] {tag}: {self.summary.get().strip()}"

    def _report(self, log_lines: int = 15) -> str:
        parts = [f"Type: {self.kind.get()}", f"Summary: {self.summary.get().strip()}"]
        if self.name.get().strip():
            parts.append(f"Submitted by: {self.name.get().strip()}")
        parts += ["", "Details:", self.details.get("1.0", "end").strip()]
        steps = self.steps.get("1.0", "end").strip()
        if self.kind.get() == self.TYPES[0] and re.search(r"\d\.\s*\S", steps):
            parts += ["", "Steps to reproduce:", steps]
        if self.include_system.get():
            parts += ["", "System:", system_info()]
        if self.include_log.get() and log_lines:
            log = self.parent.recent_log(log_lines)
            if log:
                parts += ["", "Recent activity:", log]
        return "\n".join(parts)

    def _copy_to_clipboard(self, text: str) -> None:
        self.clipboard_clear()
        self.clipboard_append(text)
        self.update()

    def copy_report(self) -> None:
        if not self._validate():
            return
        self._copy_to_clipboard(f"To: {FEEDBACK_EMAIL}\nSubject: {self._subject()}\n\n{self._report()}")
        messagebox.showinfo("Report copied",
                            f"The report is copied to the clipboard.\nPaste it into an email to {FEEDBACK_EMAIL}.",
                            parent=self)

    def send(self) -> None:
        if not self._validate():
            return
        full_report = self._report()
        subject = quote(self._subject())

        # Fit the link within mail-app limits: drop the log first, then shorten the text.
        body_text = full_report
        for log_lines in (15, 5, 0):
            body_text = self._report(log_lines)
            if len(subject) + len(quote(body_text)) < MAILTO_SAFE_LENGTH:
                break
        truncated = len(subject) + len(quote(body_text)) >= MAILTO_SAFE_LENGTH
        if truncated:
            note = "\n\n[Report shortened. Paste the full report from the clipboard here.]"
            while body_text and len(subject) + len(quote(body_text + note)) >= MAILTO_SAFE_LENGTH:
                body_text = body_text[:-50]
            body_text += note
        if truncated or body_text != full_report:
            self._copy_to_clipboard(full_report)

        url = f"mailto:{FEEDBACK_EMAIL}?subject={subject}&body={quote(body_text)}"
        try:
            open_external(url)
        except Exception:
            self._copy_to_clipboard(f"To: {FEEDBACK_EMAIL}\nSubject: {self._subject()}\n\n{full_report}")
            messagebox.showwarning(
                "Email app not available",
                f"No email app could be opened. The report is copied to the clipboard.\n"
                f"Paste it into an email to {FEEDBACK_EMAIL}.", parent=self)
            return
        message = "Your email app opened with the report. Review it, then select Send."
        if body_text != full_report:
            message += "\n\nThe email holds a shortened version. The full report is on the clipboard."
        messagebox.showinfo("Email ready", message, parent=self)
        self.destroy()


def open_external(target: str) -> None:
    """Open a file, folder or mailto link with the system's default application."""
    if sys.platform.startswith("win"):
        os.startfile(target)  # type: ignore[attr-defined]
        return
    opener = "open" if sys.platform == "darwin" else "xdg-open"
    try:
        # A frozen Linux app sets LD_LIBRARY_PATH to its own libraries; external apps must not inherit it.
        env = dict(os.environ)
        if getattr(sys, "frozen", False):
            original = env.pop("LD_LIBRARY_PATH_ORIG", None)
            if original is not None:
                env["LD_LIBRARY_PATH"] = original
            else:
                env.pop("LD_LIBRARY_PATH", None)
        subprocess.Popen([opener, target], env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except FileNotFoundError:
        if not webbrowser.open(target):
            raise OSError("no application is registered to open this item")


def system_info() -> str:
    lines = [f"App: {APP_NAME} {APP_VERSION}{' (exe)' if getattr(sys, 'frozen', False) else ''}",
             f"OS: {platform.platform()}",
             f"Python: {platform.python_version()}",
             f"PyMuPDF: {getattr(pymupdf, 'VersionBind', getattr(pymupdf, '__version__', 'unknown'))}"]
    try:
        import pymupdf4llm
        lines.append(f"pymupdf4llm: {getattr(pymupdf4llm, '__version__', 'unknown')}")
    except Exception:
        lines.append("pymupdf4llm: not available")
    return "\n".join(lines)


# =============================================================== updates
def version_tuple(text: str) -> tuple[int, ...]:
    return tuple(int(part) for part in re.findall(r"\d+", text)[:3])


def fetch_latest_release(timeout: float = 6.0) -> tuple[str, str]:
    """Ask GitHub for the latest release. Sends no document or personal data."""
    request = urllib.request.Request(
        f"https://api.github.com/repos/{REPOSITORY}/releases/latest",
        headers={"Accept": "application/vnd.github+json", "User-Agent": f"{APP_NAME}/{APP_VERSION}"})
    with urllib.request.urlopen(request, timeout=timeout) as response:
        data = json.load(response)
    return str(data.get("tag_name", "")), str(data.get("html_url") or RELEASES_URL)


# =============================================================== main window
PALETTES = {
    # Blue-gray dark theme, as in SysHealth
    "dark": {"window": "#2e3440", "card": "#3b4252", "sidebar": "#272c36", "topbar": "#272c36",
             "text": "#eceff4", "muted": "#b8c0cf", "border": "#4c566a", "input": "#434c5e",
             "select": "#4c6a8a", "accent": "#5e81ac", "accent_hover": "#6b8fbf", "accent_text": "#88c0d0",
             "on_accent": "#ffffff", "nav_active": "#3b4252", "nav_hover": "#323947",
             "button": "#434c5e", "button_hover": "#4c566a",
             "ok": "#6ccb5f", "warn": "#f2c94c", "fail": "#ff8f99", "log_bg": "#272c36", "log_fg": "#e5e9f0",
             "row_alt": "#353c4a"},
    # Light theme, as in GrepMail
    "light": {"window": "#f8f9fa", "card": "#ffffff", "sidebar": "#ffffff", "topbar": "#ffffff",
              "text": "#212529", "muted": "#6c757d", "border": "#e0e0e0", "input": "#ffffff",
              "select": "#cfe0f5", "accent": "#2b579a", "accent_hover": "#3a69b0", "accent_text": "#2b579a",
              "on_accent": "#ffffff", "nav_active": "#e8eef7", "nav_hover": "#f1f4f9",
              "button": "#eef1f5", "button_hover": "#e2e7ee",
              "ok": "#198754", "warn": "#c05c00", "fail": "#dc3545", "log_bg": "#f1f5f9", "log_fg": "#212529",
              "row_alt": "#f6f8fb"},
}

PAGES = [
    # (group, key, sidebar label, title, description, is_tool)
    ("TOOLS", "ocr", "Make searchable", "Make searchable (OCR)",
     "Adds an invisible text layer so you can search, select and copy text. Pages stay as they are.", True),
    ("TOOLS", "markdown", "Convert to Markdown", "Convert to Markdown",
     "Turns PDFs, images, Word, Excel and PowerPoint files into Markdown for AI assistants. "
     "Scanned pages are read automatically.", True),
    ("TOOLS", "reduce", "Reduce file size", "Reduce file size",
     "Makes PDFs smaller by compressing images. A file is only saved if it really gets smaller.", True),
    ("TOOLS", "tables", "Tables to Excel", "Tables to Excel",
     "Copies ruled tables (with grid lines), such as BOQs, into Excel with numbers you can calculate with.", True),
    ("LIBRARY", "library", "Search documents", "Search documents",
     "Find any word or reference number across every document in your library, offline.", False),
    ("APP", "settings", "Settings", "Settings", "Appearance, updates and library options.", False),
    ("APP", "about", "About & feedback", "About & feedback", "Version, license and how to report a problem.", False),
]
RUN_LABELS = {"ocr": "Make searchable", "markdown": "Convert to Markdown", "reduce": "Reduce size",
              "tables": "Extract tables"}


class PdfOcrApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__(className="grepall")  # Tk reports the window class as "Grepall" (matches the Linux menu entry)
        self.title(f"{APP_NAME} {APP_VERSION}")
        settings = load_settings()
        self.settings = settings
        self.palette = PALETTES[settings["theme"]]
        self.style = ttk.Style(self)
        self.style.theme_use("clam")
        base = tkfont.nametofont("TkDefaultFont")
        family = "Segoe UI" if sys.platform.startswith("win") else base.actual("family")
        self.fonts = {
            "body": (family, 10), "small": (family, 9), "bold": (family, 10, "bold"),
            "title": (family, 15, "bold"), "brand": (family, 13, "bold"), "group": (family, 8, "bold"),
            "nav": (family, 10), "mono": tkfont.nametofont("TkFixedFont"),
        }
        for name in ("TkDefaultFont", "TkTextFont", "TkMenuFont", "TkHeadingFont"):
            tkfont.nametofont(name).configure(family=family, size=10)
        self.geometry(self._fit_to_screen(settings["window_geometry"]))
        self.minsize(980, 660)

        self.files: list[Path] = []
        self.results: list[dict] = []
        self.messages: queue.Queue = queue.Queue()
        self.cancel = threading.Event()
        self.worker: threading.Thread | None = None
        self.skinned: list[tuple] = []
        self.library = DocumentLibrary()
        self.last_update_check = settings["last_update_check"]

        string_keys = ("language", "quality", "output_dir", "redact_terms", "redact_labels", "reduce_level", "theme")
        self.vars: dict[str, tk.Variable] = {}
        for key, value in settings.items():
            if key in string_keys:
                self.vars[key] = tk.StringVar(value=value)
            elif isinstance(value, bool):
                self.vars[key] = tk.BooleanVar(value=value)
        self.pages_spec = tk.StringVar()
        self.force = tk.BooleanVar(value=False)
        self.search_text = tk.StringVar()
        self.status = tk.StringVar(value="Ready")
        self.current_page = settings["last_page"] if any(p[1] == settings["last_page"] for p in PAGES) else "ocr"

        self._load_icon()
        self._build()
        self.apply_theme()
        self.show_page(self.current_page)
        self._sync_files()
        self._toggle_redaction()
        self._refresh_library_stats()
        self.after(100, self._poll)
        self.after(3000, self._maybe_check_updates)
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    # ------------------------------------------------------------ basics
    def _fit_to_screen(self, size: str) -> str:
        """Keep the window within the screen, leaving room for the taskbar."""
        width, height = (int(value) for value in size.split("x"))
        width = min(width, self.winfo_screenwidth() - 40)
        height = min(height, self.winfo_screenheight() - 90)
        return f"{max(width, 980)}x{max(height, 660)}"

    def _load_icon(self) -> None:
        self.icon_image = self.header_image = None
        path = resource_path(ICON_FILE)
        if not path.is_file():
            return
        try:
            self.icon_image = tk.PhotoImage(file=str(path))
        except tk.TclError:
            return
        self.iconphoto(True, self.icon_image)  # title bar, taskbar and pop-up windows
        factor = max(1, self.icon_image.height() // 30)  # about 30 px high in the top bar
        self.header_image = self.icon_image.subsample(factor, factor)

    def skin(self, widget, role: str):
        """Register a plain Tk widget for re-colouring when the theme changes."""
        self.skinned.append((widget, role))
        return widget

    def _role_options(self, role: str) -> dict:
        p = self.palette
        surface = {"window": p["window"], "card": p["card"], "sidebar": p["sidebar"], "topbar": p["topbar"]}
        if role in surface:
            return {"bg": surface[role]}
        if role == "line":
            return {"bg": p["border"]}
        if role == "card_border":
            return {"highlightbackground": p["border"], "highlightcolor": p["border"]}
        kind, _, where = role.partition("@")
        bg = surface.get(where or "window", p["window"])
        if kind == "text":
            return {"bg": bg, "fg": p["text"]}
        if kind == "muted":
            return {"bg": bg, "fg": p["muted"]}
        if kind == "accent":
            return {"bg": bg, "fg": p["accent_text"]}
        if kind == "primary":
            return {"bg": p["accent"], "fg": p["on_accent"], "activebackground": p["accent_hover"],
                    "activeforeground": p["on_accent"], "disabledforeground": p["muted"]}
        if kind == "button":
            return {"bg": p["button"], "fg": p["text"], "activebackground": p["button_hover"],
                    "activeforeground": p["text"], "disabledforeground": p["muted"]}
        return {}

    def style_text(self, widget) -> None:
        p = self.palette
        widget.configure(bg=p["input"], fg=p["text"], insertbackground=p["text"], relief="flat",
                         highlightthickness=1, highlightbackground=p["border"], highlightcolor=p["accent"],
                         selectbackground=p["select"], selectforeground=p["text"])

    def apply_theme(self) -> None:
        p = self.palette
        self.configure(bg=p["window"])
        s = self.style
        s.configure(".", background=p["window"], foreground=p["text"], fieldbackground=p["input"],
                    bordercolor=p["border"], lightcolor=p["border"], darkcolor=p["border"],
                    troughcolor=p["window"], selectbackground=p["select"], selectforeground=p["text"],
                    insertcolor=p["text"], focuscolor=p["accent"], font=self.fonts["body"])
        s.map(".", foreground=[("disabled", p["muted"])])
        for name, bg in (("", p["window"]), ("Card.", p["card"]), ("Side.", p["sidebar"])):
            s.configure(f"{name}TFrame", background=bg)
            s.configure(f"{name}TLabel", background=bg, foreground=p["text"])
            s.configure(f"{name}TCheckbutton", background=bg, foreground=p["text"], indicatorbackground=p["input"],
                        indicatorforeground=p["text"], upperbordercolor=p["border"], lowerbordercolor=p["border"])
            s.map(f"{name}TCheckbutton", background=[("active", bg)],
                  indicatorbackground=[("selected", p["accent"]), ("disabled", bg), ("active", p["input"])],
                  indicatorforeground=[("selected", p["on_accent"])])
            s.configure(f"{name}TRadiobutton", background=bg, foreground=p["text"], indicatorbackground=p["input"],
                        indicatorforeground=p["on_accent"])
            s.map(f"{name}TRadiobutton", background=[("active", bg)],
                  indicatorbackground=[("selected", p["accent"]), ("active", p["input"])])
        s.configure("Muted.TLabel", foreground=p["muted"])
        s.configure("Card.Muted.TLabel", background=p["card"], foreground=p["muted"])
        s.configure("Card.Head.TLabel", background=p["card"], foreground=p["text"], font=self.fonts["bold"])
        s.configure("TEntry", fieldbackground=p["input"], foreground=p["text"], padding=4)
        s.map("TEntry", fieldbackground=[("disabled", p["card"])], bordercolor=[("focus", p["accent"])],
              lightcolor=[("focus", p["accent"])])
        s.configure("TCombobox", fieldbackground=p["input"], background=p["button"], foreground=p["text"],
                    arrowcolor=p["text"], padding=3)
        s.map("TCombobox", fieldbackground=[("readonly", p["input"])], foreground=[("readonly", p["text"])],
              selectbackground=[("readonly", p["input"])], selectforeground=[("readonly", p["text"])],
              background=[("active", p["button_hover"])], bordercolor=[("focus", p["accent"])])
        self.option_add("*TCombobox*Listbox.background", p["input"])
        self.option_add("*TCombobox*Listbox.foreground", p["text"])
        self.option_add("*TCombobox*Listbox.selectBackground", p["select"])
        self.option_add("*TCombobox*Listbox.selectForeground", p["text"])
        row_height = tkfont.Font(font=self.fonts["body"]).metrics("linespace") + 8
        s.configure("Treeview", background=p["card"], fieldbackground=p["card"], foreground=p["text"],
                    rowheight=row_height, borderwidth=0)
        s.map("Treeview", background=[("selected", p["select"])], foreground=[("selected", p["text"])])
        s.configure("Treeview.Heading", background=p["window"], foreground=p["muted"], relief="flat",
                    font=self.fonts["small"], padding=(6, 4))
        s.map("Treeview.Heading", background=[("active", p["window"])])
        s.configure("Vertical.TScrollbar", background=p["button"], troughcolor=p["card"], arrowcolor=p["muted"],
                    bordercolor=p["card"], lightcolor=p["button"], darkcolor=p["button"])
        s.map("Vertical.TScrollbar", background=[("active", p["button_hover"])])
        s.configure("Horizontal.TProgressbar", background=p["accent"], troughcolor=p["input"],
                    bordercolor=p["border"], lightcolor=p["accent"], darkcolor=p["accent"])
        s.configure("Work.Horizontal.TProgressbar", background=p["accent"], troughcolor=p["input"],
                    bordercolor=p["border"], lightcolor=p["accent"], darkcolor=p["accent"], thickness=10)
        for widget, role in self.skinned:
            try:
                widget.configure(**self._role_options(role))
            except tk.TclError:
                pass
        self.style_text(self.log)
        self.log.configure(bg=p["log_bg"], fg=p["log_fg"], highlightthickness=0)
        for tree in (self.file_tree, self.result_tree, self.library_tree):
            tree.tag_configure("alt", background=p["row_alt"])
            tree.tag_configure("ok", foreground=p["ok"])
            tree.tag_configure("warn", foreground=p["warn"])
            tree.tag_configure("fail", foreground=p["fail"])
            tree.tag_configure("missing", foreground=p["muted"])
        self._update_nav()
        self._set_badge(self.status_badge_state)
        self.theme_button.configure(text="☀  Light" if self.vars["theme"].get() == "dark" else "☾  Dark")

    def toggle_theme(self) -> None:
        theme = "light" if self.vars["theme"].get() == "dark" else "dark"
        self.vars["theme"].set(theme)
        self.palette = PALETTES[theme]
        self.apply_theme()

    # ------------------------------------------------------------ small widget helpers
    def button(self, parent, text: str, command, primary: bool = False):
        widget = tk.Button(parent, text=text, command=command, relief="flat", bd=0, cursor="hand2",
                           padx=14, pady=5, highlightthickness=0, font=self.fonts["bold" if primary else "body"])
        return self.skin(widget, "primary" if primary else "button")

    def card(self, parent, title: str | None = None):
        outer = self.skin(tk.Frame(parent, highlightthickness=1), "card")
        self.skinned.append((outer, "card_border"))
        body = ttk.Frame(outer, style="Card.TFrame", padding=(14, 10, 14, 12))
        body.pack(fill="both", expand=True)
        if title:
            ttk.Label(body, text=title, style="Card.Head.TLabel").pack(anchor="w", pady=(0, 6))
        return outer, body

    def tree(self, parent, columns: list[tuple[str, str, int, bool]], height: int = 6):
        frame = ttk.Frame(parent, style="Card.TFrame")
        tree = ttk.Treeview(frame, columns=[c[0] for c in columns], show="headings", height=height,
                            selectmode="extended")
        for key, heading, width, stretch in columns:
            tree.heading(key, text=heading, anchor="w")
            tree.column(key, width=width, stretch=stretch, anchor="w")
        bar = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=bar.set)
        tree.pack(side="left", fill="both", expand=True)
        bar.pack(side="right", fill="y")
        return frame, tree

    # ------------------------------------------------------------ layout
    def _build(self) -> None:
        # Top bar
        top = self.skin(tk.Frame(self, height=52), "topbar")
        top.pack(fill="x")
        top.pack_propagate(False)
        if self.header_image:
            self.skin(tk.Label(top, image=self.header_image, bd=0), "topbar").pack(side="left", padx=(14, 0))
        self.skin(tk.Label(top, text=APP_NAME, font=self.fonts["brand"]), "accent@topbar").pack(side="left", padx=(8, 6))
        self.skin(tk.Label(top, text=APP_VERSION, font=self.fonts["small"]), "muted@topbar").pack(side="left", pady=(4, 0))
        self.skin(tk.Label(top, text="  •  " + APP_TAGLINE, font=self.fonts["small"]), "muted@topbar").pack(
            side="left", pady=(4, 0))
        self.theme_button = self.button(top, "", self.toggle_theme)
        self.theme_button.pack(side="right", padx=(6, 14))
        self.status_badge = tk.Label(top, text="", font=self.fonts["bold"], padx=12, pady=3)
        self.status_badge.pack(side="right", padx=6)
        self.status_badge_state = "ready"
        self.skin(tk.Frame(self, height=1), "line").pack(fill="x")
        self.update_bar_holder = self.skin(tk.Frame(self), "window")
        self.update_bar_holder.pack(fill="x")

        # Status bar (packed before the body so it always stays visible)
        status = self.skin(tk.Frame(self, height=30), "topbar")
        status.pack(side="bottom", fill="x")
        status.pack_propagate(False)
        self.skin(tk.Frame(self, height=1), "line").pack(side="bottom", fill="x")
        self.progress_bar = ttk.Progressbar(status, mode="determinate", length=180)
        self.progress_bar.pack(side="left", padx=(14, 10), pady=8)
        self.skin(tk.Label(status, textvariable=self.status, font=self.fonts["small"], anchor="w"),
                  "muted@topbar").pack(side="left", fill="x", expand=True)
        self.skin(tk.Label(status, text=APP_CREDIT, font=self.fonts["small"]), "muted@topbar").pack(side="right", padx=14)

        body = self.skin(tk.Frame(self), "window")
        body.pack(fill="both", expand=True)

        # Sidebar
        side = self.skin(tk.Frame(body, width=210), "sidebar")
        side.pack(side="left", fill="y")
        side.pack_propagate(False)
        self.skin(tk.Frame(body, width=1), "line").pack(side="left", fill="y")
        self.nav_items: dict[str, tuple] = {}
        group = None
        for g, key, label, *_ in PAGES:
            if g != group:
                group = g
                self.skin(tk.Label(side, text=g, font=self.fonts["group"], anchor="w"), "muted@sidebar").pack(
                    fill="x", padx=18, pady=(16, 4))
            row = self.skin(tk.Frame(side), "sidebar")
            row.pack(fill="x")
            marker = tk.Frame(row, width=3)
            marker.pack(side="left", fill="y")
            item = tk.Label(row, text=label, font=self.fonts["nav"], anchor="w", padx=15, pady=7, cursor="hand2")
            item.pack(side="left", fill="x", expand=True)
            for widget in (row, item):
                widget.bind("<Button-1>", lambda _e, k=key: self.show_page(k))
                widget.bind("<Enter>", lambda _e, k=key: self._hover_nav(k, True))
                widget.bind("<Leave>", lambda _e, k=key: self._hover_nav(k, False))
            self.nav_items[key] = (row, marker, item)

        # Content column
        content = self.skin(tk.Frame(body), "window")
        content.pack(side="left", fill="both", expand=True)
        inner = ttk.Frame(content, padding=(22, 14, 22, 10))
        inner.pack(fill="both", expand=True)
        inner.columnconfigure(0, weight=1)
        self.content = inner

        self.page_title = self.skin(tk.Label(inner, font=self.fonts["title"], anchor="w"), "text@window")
        self.page_title.grid(row=0, column=0, sticky="ew")
        self.page_desc = self.skin(tk.Label(inner, font=self.fonts["body"], anchor="w", justify="left"), "muted@window")
        self.page_desc.grid(row=1, column=0, sticky="ew", pady=(2, 10))

        self._build_files_card(inner)                       # row 2
        self.option_stack = ttk.Frame(inner)                # row 3
        self.option_stack.columnconfigure(0, weight=1)
        self.option_stack.rowconfigure(0, weight=1)
        self.page_frames: dict[str, tk.Widget] = {}
        self._build_ocr_options()
        self._build_markdown_options()
        self._build_reduce_options()
        self._build_tables_options()
        self._build_library_page()
        self._build_settings_page()
        self._build_about_page()
        self._build_action_bar(inner)                       # row 4
        self._build_results_card(inner)                     # row 5
        self._build_log(inner)                              # row 6

    def _build_files_card(self, parent) -> None:
        outer, card = self.card(parent)
        self.files_card = outer
        head = ttk.Frame(card, style="Card.TFrame")
        head.pack(fill="x")
        ttk.Label(head, text="Files", style="Card.Head.TLabel").pack(side="left")
        self.file_count = ttk.Label(head, text="", style="Card.Muted.TLabel")
        self.file_count.pack(side="left", padx=8)
        for text, command in (("Clear", self.clear_files), ("Remove", self.remove_selected),
                              ("Add folder…", self.add_folder), ("Add files…", self.add_files)):
            button = self.button(head, text, command)
            button.pack(side="right", padx=(6, 0))
        frame, self.file_tree = self.tree(card, [("name", "Name", 300, True), ("pages", "Pages", 70, False),
                                                 ("size", "Size", 90, False), ("folder", "Folder", 320, True)],
                                          height=4)
        self.file_tree.bind("<Delete>", lambda _e: self.remove_selected())
        foot = ttk.Frame(card, style="Card.TFrame")
        foot.pack(side="bottom", fill="x")  # packed before the list, so the list shrinks first
        frame.pack(fill="both", expand=True, pady=(8, 6))
        ttk.Label(foot, text="Save results to:", style="Card.TLabel").pack(side="left")
        self.output_label = ttk.Label(foot, text="", style="Card.Muted.TLabel")
        self.output_label.pack(side="left", padx=8)
        self.button(foot, "Same folder", self.reset_output).pack(side="right", padx=(6, 0))
        self.button(foot, "Change…", self.choose_output).pack(side="right", padx=(6, 0))
        ttk.Checkbutton(foot, text="Combine images into one document", variable=self.vars["combine_images"],
                        style="Card.TCheckbutton").pack(side="right", padx=16)
        self._update_output_label()

    def _common_row(self, parent, quality: bool = True, pages: bool = True) -> None:
        row = ttk.Frame(parent, style="Card.TFrame")
        row.pack(fill="x", pady=(0, 6))
        ttk.Label(row, text="Language", style="Card.TLabel").pack(side="left")
        ttk.Combobox(row, textvariable=self.vars["language"], values=list(LANGUAGES), state="readonly",
                     width=17).pack(side="left", padx=(6, 18))
        if quality:
            ttk.Label(row, text="Scan quality", style="Card.TLabel").pack(side="left")
            ttk.Combobox(row, textvariable=self.vars["quality"], values=list(RESOLUTIONS), state="readonly",
                         width=20).pack(side="left", padx=(6, 18))
        if pages:
            ttk.Label(row, text="Pages", style="Card.TLabel").pack(side="left")
            ttk.Entry(row, textvariable=self.pages_spec, width=12).pack(side="left", padx=(6, 6))
            ttk.Label(row, text="blank = all, e.g. 1-5, 9", style="Card.Muted.TLabel").pack(side="left")

    def _options_frame(self, key: str, title: str = "Options"):
        outer, body = self.card(self.option_stack, title)
        outer.grid(row=0, column=0, sticky="nsew")
        self.page_frames[key] = outer
        return body

    def _check(self, parent, text: str, key: str, **pack):
        widget = ttk.Checkbutton(parent, text=text, variable=self.vars[key], style="Card.TCheckbutton")
        widget.pack(anchor="w", **pack)
        return widget

    def _build_ocr_options(self) -> None:
        body = self._options_frame("ocr")
        self._common_row(body)
        ttk.Checkbutton(body, text="Also read pages that already contain text", variable=self.force,
                        style="Card.TCheckbutton").pack(anchor="w")
        self._check(body, "Turn sideways and upside-down pages upright", "fix_rotation")
        ttk.Label(body, text="Output: <name>_ocr.pdf, the same pages with searchable text.",
                  style="Card.Muted.TLabel").pack(anchor="w", pady=(6, 0))

    def _build_markdown_options(self) -> None:
        body = self._options_frame("markdown")
        self._common_row(body)
        row = ttk.Frame(body, style="Card.TFrame")
        row.pack(fill="x")
        for text, key in (("Page markers", "markers"), ("Remove repeating headers and footers", "strip_headers"),
                          ("Also save the searchable PDF", "save_searchable")):
            ttk.Checkbutton(row, text=text, variable=self.vars[key], style="Card.TCheckbutton").pack(
                side="left", padx=(0, 18))
        arabic = ttk.Frame(body, style="Card.TFrame")
        arabic.pack(fill="x", pady=(6, 0))
        ttk.Label(arabic, text="Arabic cleanup:", style="Card.TLabel", width=16).pack(side="left")
        for text, key in (("Digits to 0-9", "arabic_digits"), ("Remove tatweel", "arabic_tatweel"),
                          ("Remove diacritics", "arabic_diacritics"), ("Unify Alef forms", "arabic_alef")):
            ttk.Checkbutton(arabic, text=text, variable=self.vars[key], style="Card.TCheckbutton").pack(
                side="left", padx=(0, 18))
        sensitive = ttk.Frame(body, style="Card.TFrame")
        sensitive.pack(fill="x", pady=(8, 0))
        sensitive.columnconfigure(2, weight=1)
        sensitive.columnconfigure(4, weight=1)
        ttk.Label(sensitive, text="Sensitive data:", style="Card.TLabel", width=16).grid(row=0, column=0, sticky="w")
        ttk.Checkbutton(sensitive, text="Hide in Markdown", variable=self.vars["redact"], style="Card.TCheckbutton",
                        command=self._toggle_redaction).grid(row=0, column=1, columnspan=2, sticky="w")
        self.redact_pdf_check = ttk.Checkbutton(sensitive, text="Also create a redacted PDF copy",
                                                variable=self.vars["redact_pdf"], style="Card.TCheckbutton")
        self.redact_pdf_check.grid(row=0, column=3, columnspan=2, sticky="w", padx=(12, 0))
        ttk.Label(sensitive, text="Hide these words:", style="Card.TLabel").grid(row=1, column=1, sticky="w",
                                                                                padx=(0, 6), pady=(4, 0))
        self.redact_terms_entry = ttk.Entry(sensitive, textvariable=self.vars["redact_terms"], width=16)
        self.redact_terms_entry.grid(row=1, column=2, sticky="ew", pady=(4, 0))
        ttk.Label(sensitive, text="Hide the value after:", style="Card.TLabel").grid(row=1, column=3, sticky="w",
                                                                                    padx=(12, 6), pady=(4, 0))
        self.redact_labels_entry = ttk.Entry(sensitive, textvariable=self.vars["redact_labels"], width=16)
        self.redact_labels_entry.grid(row=1, column=4, sticky="ew", pady=(4, 0))
        ttk.Label(sensitive, style="Card.Muted.TLabel", wraplength=720, justify="left",
                  text="Always hides ID/Iqama, IBAN, phone, email and card numbers. Separate entries with commas. "
                       "Example: value after \"Ref, Dated\" hides the reference number and the date.")\
            .grid(row=2, column=1, columnspan=4, sticky="w", pady=(2, 0))

    def _build_reduce_options(self) -> None:
        body = self._options_frame("reduce", "Compression")
        for level, (description, _settings) in REDUCE_LEVELS.items():
            row = ttk.Frame(body, style="Card.TFrame")
            row.pack(fill="x", pady=1)
            ttk.Radiobutton(row, text=level, value=level, variable=self.vars["reduce_level"],
                            style="Card.TRadiobutton", width=11).pack(side="left")
            ttk.Label(row, text=description, style="Card.Muted.TLabel").pack(side="left")
        self._check(body, "Convert colour pages to grayscale (smaller; use for black-and-white documents)",
                    "reduce_gray", pady=(6, 0))
        ttk.Label(body, text="Output: <name>_small.pdf. Searchable text is kept. Your original file is never changed.",
                  style="Card.Muted.TLabel").pack(anchor="w", pady=(6, 0))

    def _build_tables_options(self) -> None:
        body = self._options_frame("tables")
        self._common_row(body, quality=False)
        ttk.Label(body, style="Card.Muted.TLabel", wraplength=760, justify="left",
                  text="Finds tables drawn with grid lines, in digital and scanned documents. Each table becomes a "
                       "sheet in <name>_tables.xlsx, with amounts stored as numbers. Tables without lines are not "
                       "detected. In scanned Arabic tables, check the numbers: Arabic-Indic digits are often misread.")\
            .pack(anchor="w", pady=(4, 0))

    def _build_library_page(self) -> None:
        outer, body = self.card(self.option_stack)
        outer.grid(row=0, column=0, sticky="nsew")
        self.page_frames["library"] = outer
        bar = ttk.Frame(body, style="Card.TFrame")
        bar.pack(fill="x")
        entry = ttk.Entry(bar, textvariable=self.search_text, font=self.fonts["body"])
        entry.pack(side="left", fill="x", expand=True, ipady=2)
        entry.bind("<Return>", lambda _e: self.run_search())
        self.search_entry = entry
        self.button(bar, "Search", self.run_search, primary=True).pack(side="left", padx=(8, 0))
        self.button(bar, "Show all documents", self.show_all_documents).pack(side="left", padx=(6, 0))
        ttk.Label(body, text="All words must appear on the same page. Use quotes for an exact phrase, e.g. "
                             "\"chilled water\". Arabic letter forms and digits are matched either way.",
                  style="Card.Muted.TLabel", wraplength=860, justify="left").pack(anchor="w", pady=(4, 6))
        frame, self.library_tree = self.tree(body, [("doc", "Document", 240, False), ("page", "Page", 60, False),
                                                    ("text", "Text found", 460, True), ("folder", "Folder", 220, True)],
                                             height=10)
        frame.pack(fill="both", expand=True)
        self.library_tree.bind("<Double-1>", lambda _e: self.open_library_selection())
        self.library_rows: dict[str, dict] = {}
        actions = ttk.Frame(body, style="Card.TFrame")
        actions.pack(fill="x", pady=(8, 0))
        self.button(actions, "Open document", self.open_library_selection, primary=True).pack(side="left")
        for text, command in (("Open folder", lambda: self.open_library_selection(folder=True)),
                              ("Add documents\u2026", self.add_library_documents),
                              ("Remove from library", self.remove_library_selection),
                              ("Remove missing files", self.remove_missing_documents)):
            self.button(actions, text, command).pack(side="left", padx=(6, 0))
        foot = ttk.Frame(body, style="Card.TFrame")
        foot.pack(fill="x", pady=(8, 0))
        self.library_stats = ttk.Label(foot, text="", style="Card.Muted.TLabel")
        self.library_stats.pack(side="left")
        ttk.Checkbutton(foot, text="Add documents automatically after Make searchable and Convert to Markdown",
                        variable=self.vars["library_auto"], style="Card.TCheckbutton").pack(side="right")

    def _build_settings_page(self) -> None:
        outer, body = self.card(self.option_stack, "Appearance")
        outer.grid(row=0, column=0, sticky="nsew")
        self.page_frames["settings"] = outer
        row = ttk.Frame(body, style="Card.TFrame")
        row.pack(fill="x")
        for text, value in (("Dark", "dark"), ("Light", "light")):
            ttk.Radiobutton(row, text=text, value=value, variable=self.vars["theme"], style="Card.TRadiobutton",
                            command=self._theme_from_setting).pack(side="left", padx=(0, 18))
        ttk.Label(body, text="Updates", style="Card.Head.TLabel").pack(anchor="w", pady=(14, 4))
        self._check(body, "Check for a newer version once a day (only the version number is requested)",
                    "check_updates")
        ttk.Label(body, text="Library", style="Card.Head.TLabel").pack(anchor="w", pady=(14, 4))
        self._check(body, "Add documents to the library automatically after processing", "library_auto")
        ttk.Label(body, text=f"Library file: {self.library.path}", style="Card.Muted.TLabel").pack(anchor="w", pady=(2, 0))
        ttk.Label(body, text="Reset", style="Card.Head.TLabel").pack(anchor="w", pady=(14, 4))
        row = ttk.Frame(body, style="Card.TFrame")
        row.pack(fill="x")
        self.button(row, "Reset all settings", self.reset_settings).pack(side="left")
        ttk.Label(row, text=f"   Settings file: {settings_path()}", style="Card.Muted.TLabel").pack(side="left")

    def _build_about_page(self) -> None:
        outer, body = self.card(self.option_stack)
        outer.grid(row=0, column=0, sticky="nsew")
        self.page_frames["about"] = outer
        ttk.Label(body, text=f"{APP_NAME} {APP_VERSION}", style="Card.Head.TLabel",
                  font=self.fonts["title"]).pack(anchor="w")
        ttk.Label(body, text=APP_TAGLINE, style="Card.Muted.TLabel").pack(anchor="w", pady=(0, 10))
        for line in (f"Created by {APP_CREDIT}.",
                     "Free software under the GNU Affero General Public License v3.0 (AGPL-3.0), "
                     "provided without any warranty.",
                     "All processing runs on this computer. No document is uploaded anywhere."):
            ttk.Label(body, text=line, style="Card.TLabel", wraplength=760, justify="left").pack(anchor="w", pady=1)
        row = ttk.Frame(body, style="Card.TFrame")
        row.pack(anchor="w", pady=(14, 0))
        self.button(row, "Report a bug or suggest an improvement", self.open_feedback, primary=True).pack(side="left")
        self.button(row, "Check for updates now", lambda: self._maybe_check_updates(force=True)).pack(side="left", padx=8)
        self.button(row, "Project page", lambda: open_external(f"https://github.com/{REPOSITORY}")).pack(side="left")
        ttk.Label(body, text=f"Feedback is sent through your email app to {FEEDBACK_EMAIL}.",
                  style="Card.Muted.TLabel").pack(anchor="w", pady=(8, 0))

    def _build_action_bar(self, parent) -> None:
        bar = ttk.Frame(parent)
        self.action_bar = bar
        self.run_button = self.button(bar, "", self.run_current, primary=True)
        self.run_button.pack(side="left")
        self.stop_button = self.button(bar, "Stop", self.stop)
        self.stop_button.pack(side="left", padx=8)
        self.open_out_button = self.button(bar, "Open output folder", self.open_output)
        self.open_out_button.pack(side="left")
        # Progress shown next to the buttons while a tool runs (and the result when it ends)
        self.work_panel = ttk.Frame(bar)
        self.work_panel.pack(side="left", fill="x", expand=True, padx=(16, 0))
        self.work_bar = ttk.Progressbar(self.work_panel, mode="indeterminate", length=220,
                                        style="Work.Horizontal.TProgressbar")
        self.work_label = self.skin(tk.Label(self.work_panel, text="", font=self.fonts["body"], anchor="w"),
                                    "text@window")
        self.work_started = 0.0
        self.work_determinate = False
        self.spinner_index = 0
        self.work_ticking = False
        self.work_file_label = ""

    def _build_results_card(self, parent) -> None:
        outer, card = self.card(parent)
        self.results_card = outer
        head = ttk.Frame(card, style="Card.TFrame")
        head.pack(fill="x")
        ttk.Label(head, text="Results", style="Card.Head.TLabel").pack(side="left")
        ttk.Label(head, text="double-click to open", style="Card.Muted.TLabel").pack(side="left", padx=8)
        self.button(head, "Clear", self.clear_results).pack(side="right")
        frame, self.result_tree = self.tree(card, [("tool", "Tool", 140, False), ("input", "File", 220, True),
                                                   ("output", "Saved as", 240, True), ("details", "Details", 300, True)],
                                            height=4)
        frame.pack(fill="both", expand=True, pady=(6, 0))
        self.result_tree.bind("<Double-1>", lambda _e: self.open_result())

    def _build_log(self, parent) -> None:
        holder = ttk.Frame(parent)
        holder.grid(row=6, column=0, sticky="ew", pady=(8, 0))
        head = ttk.Frame(holder)
        head.pack(fill="x")
        self.log_toggle = self.skin(tk.Label(head, text="", font=self.fonts["small"], cursor="hand2"), "muted@window")
        self.log_toggle.pack(side="left")
        self.log_toggle.bind("<Button-1>", lambda _e: self._toggle_log())
        self.feedback_link = self.skin(tk.Label(head, text="Report a problem", font=self.fonts["small"],
                                                cursor="hand2"), "accent@window")
        self.feedback_link.pack(side="right")
        self.feedback_link.bind("<Button-1>", lambda _e: self.open_feedback())
        self.log = tk.Text(holder, height=5, state="disabled", wrap="word", font=self.fonts["mono"], bd=0,
                           padx=8, pady=6)
        self.log_holder = holder
        self._show_log(self.vars["log_open"].get())

    def _toggle_log(self) -> None:
        self.vars["log_open"].set(not self.vars["log_open"].get())
        self._show_log(self.vars["log_open"].get())

    def _show_log(self, visible: bool) -> None:
        self.log_toggle.configure(text=("▾  Activity" if visible else "▸  Activity (hidden)"))
        if visible:
            self.log.pack(fill="x", pady=(4, 0))
        else:
            self.log.pack_forget()

    # ------------------------------------------------------------ navigation
    def _hover_nav(self, key: str, inside: bool) -> None:
        if key == self.current_page:
            return
        p = self.palette
        row, marker, item = self.nav_items[key]
        bg = p["nav_hover"] if inside else p["sidebar"]
        for widget in (row, marker, item):
            widget.configure(bg=bg)

    def _update_nav(self) -> None:
        p = self.palette
        for key, (row, marker, item) in self.nav_items.items():
            active = key == self.current_page
            bg = p["nav_active"] if active else p["sidebar"]
            row.configure(bg=bg)
            item.configure(bg=bg, fg=p["text"] if active else p["muted"],
                           font=self.fonts["bold"] if active else self.fonts["nav"])
            marker.configure(bg=p["accent"] if active else bg)

    def show_page(self, key: str) -> None:
        self.current_page = key
        page = next(p for p in PAGES if p[1] == key)
        self.page_title.configure(text=page[3])
        self.page_desc.configure(text=page[4])
        is_tool = page[5]
        for name, frame in self.page_frames.items():
            if name == key:
                frame.grid()
            else:
                frame.grid_remove()  # only the active page takes space
        if is_tool:
            self.files_card.grid(row=2, column=0, sticky="nsew", pady=(0, 10))
            self.option_stack.grid(row=3, column=0, sticky="ew")
            self.action_bar.grid(row=4, column=0, sticky="ew", pady=10)
            self.results_card.grid(row=5, column=0, sticky="nsew")
            self.content.rowconfigure(2, weight=1, minsize=150)
            self.content.rowconfigure(3, weight=0)
            self.content.rowconfigure(5, weight=1, minsize=120)
            self.run_button.configure(text="Working…" if self.work_started else "▶  " + RUN_LABELS[key])
            self._work_hide()
        else:
            for widget in (self.files_card, self.action_bar, self.results_card):
                widget.grid_remove()
            self.option_stack.grid(row=3, column=0, sticky="nsew")
            self.content.rowconfigure(2, weight=0, minsize=0)
            self.content.rowconfigure(3, weight=1)
            self.content.rowconfigure(5, weight=0, minsize=0)
        if key == "library":
            self._refresh_library_stats()
            self.search_entry.focus_set()
        self._update_nav()
        self._refresh_buttons()

    # ------------------------------------------------------------ settings
    def _setting_values(self) -> dict:
        settings = {key: variable.get() for key, variable in self.vars.items()}
        settings["last_update_check"] = float(self.last_update_check)
        settings["last_page"] = self.current_page
        settings["window_geometry"] = (f"{self.winfo_width()}x{self.winfo_height()}" if self.state() == "normal"
                                       else load_settings()["window_geometry"])
        return settings

    def _save_settings(self) -> None:
        save_settings(self._setting_values())

    def _theme_from_setting(self) -> None:
        self.palette = PALETTES[self.vars["theme"].get()]
        self.apply_theme()

    def reset_settings(self) -> None:
        if not messagebox.askyesno("Reset settings", "Restore all options to their default values?", parent=self):
            return
        for key, variable in self.vars.items():
            if key in DEFAULT_SETTINGS:
                variable.set(DEFAULT_SETTINGS[key])
        try:
            settings_path().unlink()
        except OSError:
            pass
        self._update_output_label()
        self._toggle_redaction()
        self._theme_from_setting()
        self._write_log("Settings restored to defaults.")

    def _toggle_redaction(self) -> None:
        state = "normal" if self.vars["redact"].get() else "disabled"
        self.redact_pdf_check.configure(state=state)
        self.redact_terms_entry.configure(state=state)
        self.redact_labels_entry.configure(state=state)

    def _cleanup_options(self) -> dict:
        return {"digits": self.vars["arabic_digits"].get(), "tatweel": self.vars["arabic_tatweel"].get(),
                "diacritics": self.vars["arabic_diacritics"].get(), "alef": self.vars["arabic_alef"].get()}

    # ------------------------------------------------------------ files
    def _sync_files(self) -> None:
        tree = self.file_tree
        tree.delete(*tree.get_children())
        for index, path in enumerate(self.files):
            try:
                size = human_size(path.stat().st_size)
            except OSError:
                size = "missing"
            pages = ""
            if path.suffix.lower() == ".pdf":
                try:
                    with pymupdf.open(path) as doc:
                        pages = str(doc.page_count)
                except Exception:
                    pages = "?"
            elif is_office(path):
                pages = OFFICE_KINDS[path.suffix.lower()]
            else:
                pages = "image"
            tree.insert("", "end", iid=str(index), values=(path.name, pages, size, str(path.parent)),
                        tags=("alt",) if index % 2 else ())
        count = len(self.files)
        self.file_count.configure(text=f"{count} file{'s' if count != 1 else ''}" if count
                                  else "PDFs, images, Word, Excel or PowerPoint files")
        self._refresh_buttons()

    def add_files(self) -> None:
        def patterns(suffixes) -> str:  # Linux file dialogs are case-sensitive
            return " ".join(f"*{suffix} *{suffix.upper()}" for suffix in suffixes)
        selected = filedialog.askopenfilenames(parent=self, title="Select files", filetypes=[
            ("All supported files", patterns(INPUT_SUFFIXES)), ("PDF files", patterns((".pdf",))),
            ("Image files", patterns(IMAGE_SUFFIXES)),
            ("Word, Excel and PowerPoint files", patterns(OFFICE_SUFFIXES)), ("All files", "*")])
        self._add_files([Path(path) for path in selected])

    def add_folder(self) -> None:
        folder = filedialog.askdirectory(parent=self, title="Select a folder of documents")
        if folder:
            found = sorted(path for path in Path(folder).iterdir()
                           if path.is_file() and path.suffix.lower() in INPUT_SUFFIXES
                           and not path.stem.endswith(OUTPUT_STEMS))
            if not found:
                messagebox.showinfo("No files found", "The selected folder contains no supported files.", parent=self)
            self._add_files(found)

    def _add_files(self, paths: list[Path]) -> None:
        for path in paths:
            if (path.suffix.lower() in INPUT_SUFFIXES and path not in self.files
                    and not path.name.startswith("~$")):  # Office lock files
                self.files.append(path)
        self._sync_files()

    def remove_selected(self) -> None:
        for index in sorted((int(i) for i in self.file_tree.selection()), reverse=True):
            del self.files[index]
        self._sync_files()

    def clear_files(self) -> None:
        self.files.clear()
        self._sync_files()

    def choose_output(self) -> None:
        folder = filedialog.askdirectory(parent=self, title="Select the output folder")
        if folder:
            self.vars["output_dir"].set(folder)
            self._update_output_label()

    def reset_output(self) -> None:
        self.vars["output_dir"].set("")
        self._update_output_label()

    def _update_output_label(self) -> None:
        folder = self.vars["output_dir"].get()
        if folder and not Path(folder).is_dir():
            self.vars["output_dir"].set("")
            folder = ""
        self.output_label.configure(text=folder or "the same folder as each file")

    def open_output(self) -> None:
        folder = self.vars["output_dir"].get()
        target = Path(folder) if folder else (self.files[0].parent if self.files else None)
        if not target or not target.exists():
            return
        try:
            open_external(str(target))
        except Exception as exc:
            messagebox.showerror("Cannot open folder", f"{target}\n\n{exc}", parent=self)

    def open_feedback(self) -> None:
        FeedbackDialog(self)

    # ------------------------------------------------------------ results
    def _add_result(self, result: dict) -> None:
        self.results.insert(0, result)
        tree = self.result_tree
        iid = str(id(result))
        result["iid"] = iid
        tree.insert("", 0, iid=iid, values=(result["tool"], result["input"], result.get("output_name", ""),
                                             result.get("details", "")), tags=(result.get("tag", "ok"),))

    def clear_results(self) -> None:
        self.results.clear()
        self.result_tree.delete(*self.result_tree.get_children())

    def open_result(self) -> None:
        selection = self.result_tree.selection()
        result = next((r for r in self.results if selection and r.get("iid") == selection[0]), None)
        if result and result.get("output") and Path(result["output"]).exists():
            try:
                open_external(result["output"])
            except Exception as exc:
                messagebox.showerror("Cannot open file", str(exc), parent=self)

    # ------------------------------------------------------------ state and messages
    def _busy(self) -> bool:
        return bool(self.worker and self.worker.is_alive())

    def _set_badge(self, state: str) -> None:
        self.status_badge_state = state
        p = self.palette
        text, color = {"ready": ("● Ready", p["ok"]), "working": ("● Working", p["warn"]),
                       "stopping": ("● Stopping", p["warn"])}[state]
        self.status_badge.configure(text=text, fg=color, bg=p["topbar"])

    def _refresh_buttons(self) -> None:
        busy = self._busy()
        tool = next(p[5] for p in PAGES if p[1] == self.current_page)
        can_run = tool and bool(self.files) and not busy
        if self.current_page == "reduce":
            can_run = can_run and any(f.suffix.lower() == ".pdf" for f in self.files)
        elif self.current_page in ("ocr", "tables"):
            can_run = can_run and any(not is_office(f) for f in self.files)
        self.run_button.configure(state="normal" if can_run else "disabled")
        self.stop_button.configure(state="normal" if busy else "disabled")
        self.open_out_button.configure(state="normal" if self.files else "disabled")
        self._set_badge("working" if busy else "ready")

    def _write_log(self, text: str) -> None:
        self.log.configure(state="normal")
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def recent_log(self, lines: int) -> str:
        content = self.log.get("1.0", "end").strip().splitlines()
        return "\n".join(content[-lines:])

    def _poll(self) -> None:
        try:
            while True:
                kind, value = self.messages.get_nowait()
                if kind == "log":
                    self._write_log(value)
                elif kind == "status":
                    self.status.set(value)
                    self._work_file(value)
                elif kind == "bar":
                    done, total = value
                    self.progress_bar.configure(maximum=max(total, 1), value=done)
                    self._work_progress(done, total)
                elif kind == "result":
                    self._add_result(value)
                elif kind == "search":
                    self._show_search_results(*value)
                elif kind == "library_changed":
                    self._refresh_library_stats()
                elif kind == "checked":
                    self.last_update_check = value
                elif kind == "update":
                    self._show_update(*value)
                elif kind == "no_update":
                    messagebox.showinfo("Updates", value, parent=self)
                elif kind == "done":
                    if self.worker:
                        self.worker.join(timeout=2)
                    self.worker = None
                    self._refresh_buttons()
                    self.status.set(value)
                    self._work_finished()
        except queue.Empty:
            pass
        self.after(100, self._poll)

    def stop(self) -> None:
        self.cancel.set()
        self._set_badge("stopping")
        self.status.set("Stopping after the current page…")

    def _on_close(self) -> None:
        if self._busy() and not messagebox.askyesno("Work in progress", "Processing is still running. Stop and close?",
                                                    parent=self):
            return
        self.cancel.set()
        self._save_settings()
        self.destroy()

    def _start(self, target) -> None:
        self.cancel.clear()
        self.progress_bar.configure(value=0)
        self.worker = threading.Thread(target=target, daemon=True)
        self.worker.start()
        self._refresh_buttons()
        if self.current_page != "library":
            self._work_begin()

    # ------------------------------------------------------------ visible progress
    SPINNER = "◐◓◑◒"

    def _work_begin(self, text: str = "Starting…") -> None:
        """Show the moving progress bar next to the buttons at once, before the first page is done."""
        import time
        if not self.work_started:
            self.work_started = time.monotonic()
            self.work_file_label = ""
        self.work_determinate = False
        self.work_bar.configure(mode="indeterminate", value=0)
        self.work_bar.pack(side="left")
        self.work_label.pack(side="left", padx=(10, 0))
        self.work_bar.start(12)
        self.run_button.configure(text="Working…")
        self.configure(cursor="watch")
        self.work_label.configure(text=f"{self.SPINNER[0]}  {text}")
        self.update_idletasks()
        if not self.work_ticking:
            self.work_ticking = True
            self._work_tick()

    def _work_file(self, status: str) -> None:
        """Restart the moving bar when the next file begins, so it never shows 100% for unfinished work."""
        match = re.match(r"File (\d+) of (\d+)", status)
        if not match or not self.work_started:
            return
        label = match.group(0)
        if label != self.work_file_label:
            self.work_file_label = label
            if self.work_determinate:
                self.work_determinate = False
                self.work_bar.configure(mode="indeterminate", value=0)
                self.work_bar.start(12)

    def _work_progress(self, done: int, total: int) -> None:
        if done <= 0 or not self.work_started:
            return
        if not self.work_determinate:
            self.work_bar.stop()
            self.work_bar.configure(mode="determinate")
            self.work_determinate = True
        self.work_bar.configure(maximum=max(total, 1), value=done)

    def _work_tick(self) -> None:
        import time
        if not self.work_started:
            self.work_ticking = False
            return
        elapsed = int(time.monotonic() - self.work_started)
        self.spinner_index = (self.spinner_index + 1) % len(self.SPINNER)
        detail = "Working"
        if self.work_determinate:
            maximum = float(self.work_bar.cget("maximum")) or 1
            value = float(self.work_bar.cget('value'))
            detail = "Saving" if value >= maximum else f"{int(100 * value / maximum)}%"
        if self.work_file_label and not self.work_file_label.endswith(" of 1"):
            detail = f"{self.work_file_label} · {detail}"
        if self.status_badge_state == "stopping":
            detail = "Stopping"
        self.work_label.configure(text=f"{self.SPINNER[self.spinner_index]}  {detail} · {elapsed // 60}:{elapsed % 60:02d}")
        self.after(250, self._work_tick)

    def _work_finished(self) -> None:
        import time
        if not self.work_started:
            return
        elapsed = int(time.monotonic() - self.work_started)
        self.work_started = 0.0
        self.work_bar.stop()
        self.work_bar.configure(mode="determinate", maximum=1, value=1)
        self.configure(cursor="")
        stopped = self.cancel.is_set()
        self.work_label.configure(text=("■  Stopped" if stopped else "✓  Finished") +
                                  f" in {elapsed // 60}:{elapsed % 60:02d}")
        self.run_button.configure(text="▶  " + RUN_LABELS[self.current_page]
                                  if self.current_page in RUN_LABELS else self.run_button.cget("text"))

    def _work_hide(self) -> None:
        if not self.work_started:
            self.work_bar.pack_forget()
            self.work_label.pack_forget()

    # ------------------------------------------------------------ updates
    def _maybe_check_updates(self, force: bool = False) -> None:
        if not force and (not self.vars["check_updates"].get()
                          or time.time() - self.last_update_check < UPDATE_CHECK_INTERVAL):
            return

        def worker() -> None:
            try:
                tag, url = fetch_latest_release()
            except Exception:
                if force:
                    self.messages.put(("no_update", "Could not reach GitHub. Check your internet connection."))
                return  # offline or blocked: try again at the next start
            self.messages.put(("checked", time.time()))
            if tag and version_tuple(tag) > version_tuple(APP_VERSION):
                self.messages.put(("update", (tag.lstrip("vV"), url)))
            elif force:
                self.messages.put(("no_update", f"You have the latest version ({APP_VERSION})."))

        threading.Thread(target=worker, daemon=True).start()

    def _show_update(self, version: str, url: str) -> None:
        for child in self.update_bar_holder.winfo_children():
            child.destroy()
        p = self.palette
        bar = tk.Frame(self.update_bar_holder, bg=p["accent"])
        bar.pack(fill="x")
        tk.Label(bar, text=f"   GrepAll {version} is available.", bg=p["accent"], fg=p["on_accent"],
                 font=self.fonts["bold"]).pack(side="left", pady=6)
        tk.Button(bar, text="Download", command=lambda: open_external(url), relief="flat", bd=0, cursor="hand2",
                  bg=p["on_accent"], fg=p["accent"], padx=12).pack(side="left", padx=10)
        tk.Button(bar, text="Dismiss", command=bar.destroy, relief="flat", bd=0, cursor="hand2",
                  bg=p["accent"], fg=p["on_accent"], activebackground=p["accent_hover"]).pack(side="right", padx=10)
        self._write_log(f"Update available: GrepAll {version} ({url})")

    # ------------------------------------------------------------ running tools
    def _plan_jobs(self, pdf_only: bool = False, office: bool = False) -> list[tuple[str, list[Path], Path]] | None:
        """Build (title, inputs, output folder) jobs. Returns None if the user cancels."""
        output_dir = self.vars["output_dir"].get()
        jobs = []
        files = [f for f in self.files if (not pdf_only or f.suffix.lower() == ".pdf") and (office or not is_office(f))]
        images = [path for path in files if is_image(path)]
        combine = self.vars["combine_images"].get() and len(images) >= 2
        titles: set[str] = set()
        for path in files:
            if combine and is_image(path):
                continue
            title = path.stem
            if title in titles:  # e.g. report.pdf and report.jpg in the same run
                title = f"{path.stem}_{path.suffix.lstrip('.').lower()}"
            titles.add(title)
            jobs.append((title, [path], Path(output_dir) if output_dir else path.parent))
        if combine:
            initial_dir = output_dir or str(images[0].parent)
            chosen = filedialog.asksaveasfilename(
                parent=self, title="Name for the combined images", initialdir=initial_dir,
                initialfile="combined_images.pdf", defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")])
            if not chosen:
                return None
            chosen_path = Path(chosen)
            title = re.sub(r"_ocr$", "", chosen_path.stem)
            jobs.append((title, images, chosen_path.parent))
        return jobs

    def run_current(self) -> None:
        if self._busy() or not self.files:
            return
        self._work_begin("Preparing…")
        try:
            self._run_current()
        finally:
            if not self._busy() and self.work_started:   # stopped before starting (error or cancelled dialog)
                self.work_started = 0.0
                self.work_bar.stop()
                self.configure(cursor="")
                self.run_button.configure(text="▶  " + RUN_LABELS[self.current_page])
                self.work_bar.pack_forget()
                self.work_label.pack_forget()

    def _run_current(self) -> None:
        tool = self.current_page
        language = LANGUAGES[self.vars["language"].get()]
        page_spec = self.pages_spec.get().strip() or None
        try:
            parse_pages(page_spec, 10**6)  # syntax check only
        except ValueError:
            messagebox.showerror("Invalid page selection", "Enter pages as numbers or ranges, for example: 1-5, 9",
                                 parent=self)
            return
        tessdata = ""
        if tool != "reduce":
            try:
                tessdata = find_tessdata()
                check_languages(tessdata, language)
            except Exception as exc:
                messagebox.showerror("OCR data not ready", str(exc), parent=self)
                return
        jobs = self._plan_jobs(pdf_only=tool == "reduce", office=tool == "markdown")
        if not jobs:
            return
        skipped = sum(1 for f in self.files if is_office(f)) if tool != "markdown" else 0
        if skipped:
            self._write_log(f"Skipped {skipped} Word/Excel/PowerPoint file(s): only Convert to Markdown reads them.")
        self._save_settings()
        options = {
            "tool": tool, "language": language, "tessdata": tessdata, "page_spec": page_spec,
            "dpi": RESOLUTIONS[self.vars["quality"].get()], "force": self.force.get(),
            "fix_rotation": self.vars["fix_rotation"].get(), "library_auto": self.vars["library_auto"].get(),
            "markers": self.vars["markers"].get(), "strip_headers": self.vars["strip_headers"].get(),
            "save_searchable": self.vars["save_searchable"].get(), "cleanup": self._cleanup_options(),
            "reduce_level": self.vars["reduce_level"].get(), "reduce_gray": self.vars["reduce_gray"].get(),
            "redaction": None,
        }
        if tool == "markdown" and self.vars["redact"].get():
            terms = parse_terms(self.vars["redact_terms"].get())
            labels = parse_labels(self.vars["redact_labels"].get())
            if any(options["cleanup"].values()):
                terms += [clean_arabic(t, **options["cleanup"]) for t in terms]
                labels += [clean_arabic(label, **options["cleanup"]) for label in labels]
            options["redaction"] = {"terms": sorted(set(terms)), "labels": sorted(set(labels)),
                                    "pdf": self.vars["redact_pdf"].get()}
        self._write_log(f"--- {RUN_LABELS[tool]} started ({len(jobs)} file{'s' if len(jobs) != 1 else ''}) ---")
        self._start(lambda: self._tool_worker(jobs, options))

    def _prepare_text_pdf(self, inputs: list[Path], workdir: Path, options: dict, label: str, job_text: str):
        """A PDF of the selected pages in which every page has text (scanned pages are read with OCR).
        Returns (path, original page numbers) or None when no selected page exists."""
        put = self.messages.put
        source = inputs[0]
        if is_image(source):
            source = workdir / "images.pdf"
            images_to_pdf(inputs, source)
        with pymupdf.open(source) as probe:
            count = probe.page_count
        try:
            pages = parse_pages(options["page_spec"], count)
        except ValueError:
            put(("log", f"  Skipped: no selected pages exist in this {count}-page file."))
            return None
        text_pdf = workdir / "text.pdf"

        def progress(current: int, total: int, page_number: int, status: str) -> None:
            put(("bar", (current, total)))
            put(("status", f"{job_text}: {label}, page {page_number} ({current} of {total})"))
            if "upright" in status or "large page" in status:
                put(("log", f"  Page {page_number}: {status}"))
            if self.cancel.is_set():
                raise Cancelled()

        ocr_count, copied = ocr_pdf(source, text_pdf, pages, options["language"], options["dpi"], options["tessdata"],
                                    options["force"] if options["tool"] == "ocr" else False, progress,
                                    options["fix_rotation"])
        return text_pdf, [p + 1 for p in pages], ocr_count, copied

    def _library_add(self, path: Path, texts: list[str]) -> None:
        try:
            self.library.add(path, texts)
            self.messages.put(("log", f"  Added to library: {path.name}"))
            self.messages.put(("library_changed", None))
        except Exception as exc:
            self.messages.put(("log", f"  Could not add to library: {exc}"))

    def _tool_worker(self, jobs, options: dict) -> None:
        put = self.messages.put
        tool = options["tool"]
        completed = 0
        for job_index, (title, inputs, out_dir) in enumerate(jobs, 1):
            if self.cancel.is_set():
                break
            label = inputs[0].name if len(inputs) == 1 else f"{len(inputs)} images ({title})"
            job_text = f"File {job_index} of {len(jobs)}"
            put(("log", str(inputs[0]) if len(inputs) == 1 else f"{len(inputs)} images combined as {title}"))
            started = time.time()
            workdir = Path(tempfile.mkdtemp(prefix="grepall-"))
            try:
                out_dir.mkdir(parents=True, exist_ok=True)
                if tool == "reduce":
                    put(("status", f"{job_text}: reducing {label}"))
                    put(("bar", (0, 1)))
                    output = out_dir / f"{title}_small.pdf"
                    before, after, saved = reduce_pdf(inputs[0], output, options["reduce_level"], options["reduce_gray"])
                    put(("bar", (1, 1)))
                    if saved:
                        change = min(99, round(100 * (1 - after / before)))
                        change_text = f"{change}% smaller" if change >= 1 else "less than 1% smaller"
                        details = f"{human_size(before)} \u2192 {human_size(after)}  ({change_text})"
                        put(("log", f"  Saved {output.name}: {details}"))
                        put(("result", {"tool": "Reduce size", "input": label, "output": str(output),
                                        "output_name": output.name, "details": details, "tag": "ok"}))
                    else:
                        details = f"Already compact ({human_size(before)}); nothing saved"
                        put(("log", f"  {details}"))
                        put(("result", {"tool": "Reduce size", "input": label, "details": details, "tag": "warn"}))
                    completed += 1
                    continue

                if is_office(inputs[0]):
                    put(("status", f"{job_text}: converting {label}"))
                    put(("bar", (0, 1)))
                    md_path = out_dir / f"{title}.md"
                    redaction = options["redaction"]
                    texts, counts, details = office_to_markdown(inputs[0], md_path, title, options["cleanup"],
                                                                redaction and redaction["terms"],
                                                                redaction["labels"] if redaction else ())
                    put(("bar", (1, 1)))
                    put(("log", f"  Saved {md_path.name}: {details}"))
                    if redaction is not None:
                        summary = ", ".join(f"{count} {name}" for name, count in sorted(counts.items()))
                        put(("log", f"  Hidden in Markdown: {summary or 'nothing found'}"))
                        details += f"; hidden: {sum(counts.values())}"
                        if redaction["pdf"] and not options.get("office_pdf_note"):
                            options["office_pdf_note"] = True
                            put(("log", "  Note: redacted PDF copies are made for PDFs and images only."))
                    put(("result", {"tool": "Markdown", "input": label, "output": str(md_path),
                                    "output_name": md_path.name, "details": details,
                                    "tag": "warn" if "missing" in details else "ok"}))
                    if options["library_auto"]:
                        self._library_add(inputs[0], texts)
                    completed += 1
                    continue

                prepared = self._prepare_text_pdf(inputs, workdir, options, label, job_text)
                if prepared is None:
                    continue
                text_pdf, page_numbers, ocr_count, copied = prepared
                read = f"{ocr_count} page(s) read with OCR, {copied} with existing text"

                if tool == "ocr":
                    output = out_dir / f"{title}_ocr.pdf"
                    shutil.move(str(text_pdf), str(output))
                    put(("log", f"  Saved {output.name}: {read}, {time.time() - started:.0f}s"))
                    put(("result", {"tool": "Make searchable", "input": label, "output": str(output),
                                    "output_name": output.name, "details": read, "tag": "ok"}))
                    if options["library_auto"]:
                        self._library_add(output, page_texts(output))

                elif tool == "markdown":
                    md_path = out_dir / f"{title}.md"
                    redaction = options["redaction"]
                    warning, counts = to_markdown(text_pdf, md_path, title, page_numbers, options["markers"],
                                                  options["strip_headers"], options["cleanup"],
                                                  redaction and redaction["terms"],
                                                  redaction["labels"] if redaction else ())
                    if warning:
                        put(("log", f"  Note: {warning}"))
                    details = read
                    put(("log", f"  Saved {md_path.name}"))
                    if redaction is not None:
                        summary = ", ".join(f"{count} {name}" for name, count in sorted(counts.items()))
                        put(("log", f"  Hidden in Markdown: {summary or 'nothing found'}"))
                        details += f"; hidden: {sum(counts.values())}"
                        if redaction["pdf"]:
                            redacted = out_dir / f"{title}_redacted.pdf"
                            areas = redact_pdf(text_pdf, redacted, redaction["terms"], redaction["labels"])
                            put(("log", f"  Saved {redacted.name}: {areas} area(s) blacked out"))
                            put(("result", {"tool": "Redacted copy", "input": label, "output": str(redacted),
                                            "output_name": redacted.name, "details": f"{areas} area(s) blacked out",
                                            "tag": "ok"}))
                    searchable = None
                    if options["save_searchable"]:
                        searchable = out_dir / f"{title}_ocr.pdf"
                        shutil.copy2(text_pdf, searchable)
                        put(("log", f"  Saved {searchable.name}"))
                    put(("result", {"tool": "Markdown", "input": label, "output": str(md_path),
                                    "output_name": md_path.name, "details": details, "tag": "ok"}))
                    if options["library_auto"]:
                        target = searchable or (inputs[0] if len(inputs) == 1 else None)
                        if target is not None:
                            self._library_add(target, page_texts(text_pdf))

                elif tool == "tables":
                    def progress(current: int, total: int, page_number: int, status: str) -> None:
                        put(("bar", (current, total)))
                        put(("status", f"{job_text}: finding tables in {label}, page {current} of {total}"))
                        if self.cancel.is_set():
                            raise Cancelled()
                    with pymupdf.open(text_pdf) as probe:
                        indexes = list(range(probe.page_count))
                    tables = extract_tables(text_pdf, indexes, options["language"], options["tessdata"], progress)
                    for table in tables:
                        table["page"] = page_numbers[table["page"] - 1]
                    if tables:
                        output = out_dir / f"{title}_tables.xlsx"
                        write_tables_xlsx(tables, output, title)
                        pages_with = len({t["page"] for t in tables})
                        details = f"{len(tables)} table(s) on {pages_with} page(s)"
                        put(("log", f"  Saved {output.name}: {details}"))
                        put(("result", {"tool": "Tables to Excel", "input": label, "output": str(output),
                                        "output_name": output.name, "details": details, "tag": "ok"}))
                    else:
                        put(("log", "  No ruled tables found (tables need visible grid lines)."))
                        put(("result", {"tool": "Tables to Excel", "input": label,
                                        "details": "No ruled tables found", "tag": "warn"}))
                completed += 1
            except Cancelled:
                put(("log", "  Stopped. This file was not saved."))
                break
            except Exception as exc:
                put(("log", f"  Failed: {exc}"))
                put(("result", {"tool": RUN_LABELS[tool], "input": label, "details": f"Failed: {exc}", "tag": "fail"}))
            finally:
                shutil.rmtree(workdir, ignore_errors=True)
        if options["redaction"] is not None and completed:
            put(("log", f"  Note: {REDACTION_NOTE}"))
        result = (f"Stopped. {completed} file(s) done." if self.cancel.is_set()
                  else f"{RUN_LABELS[tool]} finished: {completed} of {len(jobs)} file(s).")
        put(("log", result))
        put(("done", result))

    # ------------------------------------------------------------ library
    def _refresh_library_stats(self) -> None:
        try:
            documents, pages = self.library.stats()
        except Exception:
            documents = pages = 0
        self.library_stats.configure(text=f"{documents} document(s), {pages} page(s) in the library")

    def run_search(self) -> None:
        query = self.search_text.get().strip()
        if not query:
            return
        try:
            results = self.library.search(query)
        except Exception as exc:
            messagebox.showerror("Search failed", str(exc), parent=self)
            return
        self._show_search_results(results, query)

    def _show_search_results(self, results: list[dict], query: str) -> None:
        tree = self.library_tree
        tree.delete(*tree.get_children())
        self.library_rows = {}
        tree.heading("page", text="Page")
        tree.heading("text", text="Text found")
        for index, row in enumerate(results):
            iid = f"r{index}"
            self.library_rows[iid] = row
            tags = ["alt"] if index % 2 else []
            if not row["exists"]:
                tags.append("missing")
            tree.insert("", "end", iid=iid, tags=tuple(tags), values=(
                row["title"] + ("" if row["exists"] else "  (missing)"), self._page_label(row), row["snippet"],
                str(Path(row["path"]).parent)))
        found = len(results)
        self.status.set(f"{found} page(s) found for “{query}”" + (" (first 500 shown)" if found >= 500 else ""))

    @staticmethod
    def _page_label(row: dict) -> str:
        kind = OFFICE_KINDS.get(Path(row["path"]).suffix.lower())
        return {"Excel": f"Sheet {row['page']}", "PowerPoint": f"Slide {row['page']}", "Word": "—"}.get(kind,
                                                                                                    str(row["page"]))

    def show_all_documents(self) -> None:
        tree = self.library_tree
        tree.delete(*tree.get_children())
        self.library_rows = {}
        tree.heading("page", text="Pages")
        tree.heading("text", text="Added")
        for index, doc in enumerate(self.library.documents()):
            iid = f"d{index}"
            self.library_rows[iid] = {**doc, "page": 1, "doc_id": doc["id"]}
            tags = ["alt"] if index % 2 else []
            if not doc["exists"]:
                tags.append("missing")
            tree.insert("", "end", iid=iid, tags=tuple(tags), values=(
                doc["title"] + ("" if doc["exists"] else "  (missing)"), doc["pages"],
                time.strftime("%Y-%m-%d %H:%M", time.localtime(doc["added"])), str(Path(doc["path"]).parent)))
        self._refresh_library_stats()

    def open_library_selection(self, folder: bool = False) -> None:
        selection = self.library_tree.selection()
        if not selection:
            return
        row = self.library_rows.get(selection[0])
        if not row:
            return
        path = Path(row["path"])
        target = path.parent if folder else path
        if not target.exists():
            messagebox.showwarning("File not found", f"{path}\n\nThe file was moved or deleted. Use "
                                   "\"Remove missing files\" to clean up the library.", parent=self)
            return
        try:
            open_external(str(target))
            if not folder:
                self.status.set(f"Opened {path.name}. The text was found on page {row['page']}.")
        except Exception as exc:
            messagebox.showerror("Cannot open", str(exc), parent=self)

    def remove_library_selection(self) -> None:
        ids = sorted({self.library_rows[i]["doc_id"] for i in self.library_tree.selection() if i in self.library_rows})
        if not ids or not messagebox.askyesno(
                "Remove from library", f"Remove {len(ids)} document(s) from the library?\n\n"
                                       "The files themselves are not deleted.", parent=self):
            return
        self.library.remove(ids)
        self.library_tree.delete(*self.library_tree.get_children())
        self._refresh_library_stats()
        self._write_log(f"Removed {len(ids)} document(s) from the library.")

    def remove_missing_documents(self) -> None:
        removed = self.library.remove_missing()
        self._refresh_library_stats()
        self.library_tree.delete(*self.library_tree.get_children())
        messagebox.showinfo("Library", f"{removed} missing document(s) removed from the library.", parent=self)

    def add_library_documents(self) -> None:
        if self._busy():
            return
        def patterns(suffixes) -> str:
            return " ".join(f"*{suffix} *{suffix.upper()}" for suffix in suffixes)
        selected = filedialog.askopenfilenames(parent=self, title="Add documents to the library", filetypes=[
            ("All supported files", patterns(INPUT_SUFFIXES)), ("All files", "*")])
        paths = [Path(p) for p in selected if Path(p).suffix.lower() in INPUT_SUFFIXES]
        if not paths:
            return
        language = LANGUAGES[self.vars["language"].get()]
        try:
            tessdata = find_tessdata()
            check_languages(tessdata, language)
        except Exception as exc:
            messagebox.showerror("OCR data not ready", str(exc), parent=self)
            return
        options = {"tool": "library", "language": language, "tessdata": tessdata, "page_spec": None,
                   "dpi": RESOLUTIONS[self.vars["quality"].get()], "force": False,
                   "fix_rotation": self.vars["fix_rotation"].get()}
        self._write_log(f"--- Adding {len(paths)} document(s) to the library ---")
        self._start(lambda: self._library_worker(paths, options))

    def _library_worker(self, paths: list[Path], options: dict) -> None:
        put = self.messages.put
        added = 0
        for index, path in enumerate(paths, 1):
            if self.cancel.is_set():
                break
            workdir = Path(tempfile.mkdtemp(prefix="grepall-"))
            try:
                if is_office(path):
                    put(("status", f"Document {index} of {len(paths)}: reading {path.name}"))
                    self._library_add(path, office_texts(path))
                    added += 1
                    continue
                prepared = self._prepare_text_pdf([path], workdir, options, path.name, f"Document {index} of {len(paths)}")
                if prepared:
                    self._library_add(path, page_texts(prepared[0]))
                    added += 1
            except Cancelled:
                break
            except Exception as exc:
                put(("log", f"  {path.name}: failed ({exc})"))
            finally:
                shutil.rmtree(workdir, ignore_errors=True)
        result = f"Library: {added} of {len(paths)} document(s) added."
        put(("log", result))
        put(("library_changed", None))
        put(("done", result))


def main() -> None:
    # A windowed .exe has no console: give print() somewhere harmless to write.
    if sys.stdout is None:
        sys.stdout = open(os.devnull, "w", encoding="utf-8")
    if sys.stderr is None:
        sys.stderr = open(os.devnull, "w", encoding="utf-8")
    app = PdfOcrApp()
    app.mainloop()


if __name__ == "__main__":
    main()
